"""Stage 2C: standard Measurement/Evidence artifact contract."""
from __future__ import annotations

import csv
import json
import re
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

_SRC = Path(__file__).resolve().parents[1] / "src"
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from mouse_click_stability_analyzer import app as legacy_app
from mouse_click_stability_analyzer import report_v2
from mouse_click_stability_analyzer.artifact_contract import (
    FACTUAL_EVENT_FIELDS,
    FORBIDDEN_CSV_COLUMNS,
    STANDARD_OPTIONAL_KEYS,
    STANDARD_REQUIRED_KEYS,
    scan_csv_headers,
    scan_forbidden_semantics,
)
from mouse_click_stability_analyzer.chart_domain import (
    DOMAIN_COUNT,
    DOMAIN_ELAPSED,
    DOMAIN_HOLD,
    DOMAIN_INTERVAL,
    DOMAIN_SIGNED,
    domain_axis_bounds,
)
from mouse_click_stability_analyzer.completed_result import (
    RUN_STATUS_COMPLETE,
    build_completed_result_model,
)
from mouse_click_stability_analyzer.schema import artifact_names, legacy_artifact_names, report_package_dir


def _events(n: int = 8, interval_ms: float = 100.0, hold_ms: float = 40.0) -> list[dict]:
    presses = [i * interval_ms for i in range(n)]
    out = []
    for i, press in enumerate(presses):
        out.append({
            "index": i + 1,
            "click_index": i + 1,
            "button": "Left",
            "accepted": True,
            "press_elapsed_ms": press - presses[0],
            "press_interval_ms": None if i == 0 else interval_ms,
            "hold_ms": hold_ms,
            "capture_source": "Windows Raw Input",
            "device_source": r"\\?\HID#VID_0000",
            "click_class": "MISS_OR_SLOW",
            "analysis_status": "WARN",
            "issue_tags": "INTERVAL_JITTER_WARN",
            "expected_click_index": i + 1,
            "expected_time_ms": i * interval_ms,
            "drift_ms": 0.4,
        })
    return out


def _fixture_settings() -> dict:
    return {
        "profile_name": "Fixture_Custom",
        "manual_mode": False,
        "focus_button": "Left",
        "button": "Left",
        "fixture_hz": 10.0,
        "fixture_reference_hz": 10.0,
        "duration_sec": 2.0,
        "capture_mode": "Windows Raw Input",
        "language": "en",
        "fixture_timing_tolerance_ms": 2.0,
        "chatter_threshold_ms": 8.0,
        "duplicate_click_threshold_ms": 30.0,
        "interval_cv_limit_pct": 12.0,
    }


class StandardBundleContractTests(unittest.TestCase):
    def test_standard_bundle_filenames_and_verifier(self):
        events = _events()
        settings = _fixture_settings()
        model = build_completed_result_model(
            locked=settings, events=events, run_id="std1", target="DUT", operator="op",
        )
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            paths = legacy_app.write_reports(root, "DUT", "op", settings, events, "std1", completed_model=model)
            names = artifact_names("std1")
            self.assertEqual(set(names), {"evidence_json", "all_click_log_csv", "engineering_html", "engineering_xlsx"})
            verification = legacy_app._verify_artifacts(paths)
            self.assertTrue(verification["ok"])
            self.assertEqual(
                verification["required"],
                ["evidence_json", "all_click_log_csv", "engineering_html", "engineering_xlsx"],
            )
            self.assertNotIn("optional", verification)
            self.assertNotIn("xlsx_skipped", verification)
            package = root / "mcsa_std1"
            for fname in names.values():
                self.assertTrue((package / fname).is_file(), fname)
            for fname in legacy_artifact_names("std1").values():
                self.assertFalse((package / fname).exists(), fname)
                self.assertFalse((root / fname).exists(), fname)
            evidence = json.loads(Path(paths["evidence_json"]).read_text(encoding="utf-8"))
            self.assertEqual(
                evidence["report_artifacts_verified"]["required"],
                ["evidence_json", "events_csv", "report_html", "report_xlsx"],
            )
            self.assertTrue(evidence["report_artifacts_verified"]["ok"])

    def test_forbidden_semantics_absent_from_standard_json(self):
        events = _events()
        settings = _fixture_settings()
        model = build_completed_result_model(
            locked=settings, events=events, run_id="std2", target="DUT", operator="op",
        )
        with tempfile.TemporaryDirectory() as tmp:
            paths = legacy_app.write_reports(Path(tmp), "DUT", "op", settings, events, "std2", completed_model=model)
            evidence = json.loads(Path(paths["evidence_json"]).read_text(encoding="utf-8"))
        hits = scan_forbidden_semantics(evidence, source="evidence")
        self.assertEqual(hits, [])
        self.assertEqual(evidence["run_status"], RUN_STATUS_COMPLETE)
        self.assertEqual(evidence["schema"], "MCSA_EVIDENCE_V1")
        self.assertNotIn("pass", evidence)
        self.assertNotIn("stability_score", evidence)
        self.assertNotIn("dut_score", evidence)
        self.assertNotIn("setup_score", evidence)
        self.assertNotIn("issue_interpretation_map", evidence)
        self.assertNotIn("next_steps", evidence)
        self.assertNotIn("chatter_threshold_ms", evidence["settings"])
        self.assertEqual(evidence["settings"]["fixture_timing_tolerance_ms"], 2.0)
        self.assertIsInstance(evidence.get("host_information"), dict)
        self.assertIn("host_press_count", evidence["measurements"])

    def test_factual_events_csv_whitelist(self):
        events = _events()
        settings = _fixture_settings()
        model = build_completed_result_model(
            locked=settings, events=events, run_id="std3", target="DUT", operator="op",
        )
        with tempfile.TemporaryDirectory() as tmp:
            paths = legacy_app.write_reports(Path(tmp), "DUT", "op", settings, events, "std3", completed_model=model)
            with Path(paths["all_click_log_csv"]).open(encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                headers = list(reader.fieldnames or [])
                rows = list(reader)
        self.assertEqual(headers, list(FACTUAL_EVENT_FIELDS))
        self.assertEqual(scan_csv_headers(headers), [])
        for name in FORBIDDEN_CSV_COLUMNS:
            self.assertNotIn(name, headers)
        self.assertEqual(len(rows), len(events))
        for raw, row in zip(events, rows):
            self.assertEqual(row["index"], str(raw["index"]))
            self.assertEqual(row["button"], raw["button"])
            self.assertEqual(row["accepted"], "True")
            self.assertEqual(float(row["press_elapsed_ms"]), raw["press_elapsed_ms"])
            self.assertNotIn("expected_click_index", row)
            self.assertNotIn("click_class", row)

    def test_missing_any_required_artifact_fails_verification(self):
        events = _events()
        settings = _fixture_settings()
        with tempfile.TemporaryDirectory() as tmp:
            paths = legacy_app.write_reports(Path(tmp), "DUT", "op", settings, events, "optx")
            for key in STANDARD_REQUIRED_KEYS:
                blank = dict(paths)
                blank[key] = ""
                verification = legacy_app._verify_artifacts(blank)
                self.assertFalse(verification["ok"], key)
                self.assertEqual(verification["missing"], [key])
                Path(paths[key]).rename(paths[key] + ".gone")
                verification = legacy_app._verify_artifacts(paths)
                self.assertFalse(verification["ok"], key)
                self.assertIn(key, verification["missing"])
                Path(paths[key] + ".gone").rename(paths[key])
        self.assertEqual(STANDARD_OPTIONAL_KEYS, ())

    def test_xlsx_writer_failure_means_generation_failed(self):
        events = _events()
        settings = _fixture_settings()
        with tempfile.TemporaryDirectory() as tmp:
            with mock.patch.object(legacy_app, "write_report_xlsx", return_value=False):
                with self.assertRaises(RuntimeError) as ctx:
                    legacy_app.write_reports(Path(tmp), "DUT", "op", settings, events, "noxl")
            self.assertIn("engineering_xlsx", str(ctx.exception))


class ReportPackageTests(unittest.TestCase):
    def _write(self, root: Path, rid: str) -> dict:
        events = _events()
        settings = _fixture_settings()
        model = build_completed_result_model(
            locked=settings, events=events, run_id=rid, target="DUT", operator="op",
        )
        return legacy_app.write_reports(root, "DUT", "op", settings, events, rid, completed_model=model)

    def test_one_package_dir_with_exactly_four_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            paths = self._write(root, "pkg1")
            children = list(root.iterdir())
            self.assertEqual([p.name for p in children], ["mcsa_pkg1"])
            self.assertTrue(children[0].is_dir())
            files = sorted(p.name for p in children[0].iterdir())
            self.assertEqual(files, sorted(artifact_names("pkg1").values()))
            self.assertEqual(sorted(files), sorted([
                "mcsa_pkg1_events.csv", "mcsa_pkg1_evidence.json",
                "mcsa_pkg1_report.html", "mcsa_pkg1_report.xlsx",
            ]))
            for value in paths.values():
                self.assertEqual(Path(value).parent, report_package_dir(root, "pkg1"))
            self.assertFalse([p for p in children[0].iterdir() if p.is_dir()])

    def test_evidence_artifact_references_are_package_relative(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            paths = self._write(root, "rel1")
            raw = Path(paths["evidence_json"]).read_text(encoding="utf-8")
            evidence = json.loads(raw)
        self.assertEqual(evidence["artifacts"], {
            "evidence_json": "mcsa_rel1_evidence.json",
            "events_csv": "mcsa_rel1_events.csv",
            "report_html": "mcsa_rel1_report.html",
            "report_xlsx": "mcsa_rel1_report.xlsx",
        })
        for entry in evidence["report_artifacts_verified"]["status"].values():
            self.assertTrue(entry["exists"])
            self.assertNotIn("path", entry)
            self.assertEqual(Path(entry["file"]).name, entry["file"])
        self.assertNotIn(str(root), raw)
        self.assertNotIn(str(root).replace("\\", "\\\\"), raw)
        self.assertNotRegex(raw, r'"[A-Za-z]:\\\\')

    def test_package_portable_across_roots(self):
        with tempfile.TemporaryDirectory() as tmp_a, tempfile.TemporaryDirectory() as tmp_b:
            paths_a = self._write(Path(tmp_a) / "reports", "port1")
            moved = Path(tmp_b) / "elsewhere" / "mcsa_port1"
            shutil.copytree(Path(paths_a["evidence_json"]).parent, moved)
            evidence = json.loads((moved / "mcsa_port1_evidence.json").read_text(encoding="utf-8"))
            for name in evidence["artifacts"].values():
                self.assertTrue((moved / name).is_file(), name)
            paths_b = self._write(Path(tmp_b) / "other_root", "port1")
            evidence_b = json.loads(Path(paths_b["evidence_json"]).read_text(encoding="utf-8"))
        self.assertEqual(evidence["artifacts"], evidence_b["artifacts"])

    def test_report_generated_at_not_completed_at(self):
        with tempfile.TemporaryDirectory() as tmp:
            paths = self._write(Path(tmp), "ts1")
            evidence = json.loads(Path(paths["evidence_json"]).read_text(encoding="utf-8"))
        self.assertIn("report_generated_at", evidence)
        self.assertTrue(evidence["report_generated_at"])
        self.assertNotIn("completed_at", evidence)
        self.assertNotIn("created_at", evidence)
        self.assertIn("elapsed_s", evidence)

    def test_html_and_xlsx_use_analysis_scope_label(self):
        with tempfile.TemporaryDirectory() as tmp:
            paths = self._write(Path(tmp), "term1")
            html_text = Path(paths["engineering_html"]).read_text(encoding="utf-8")
            from openpyxl import load_workbook
            wb = load_workbook(paths["engineering_xlsx"], read_only=True)
            cells = [
                str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None
            ]
            wb.close()
        self.assertIn("Analysis scope", html_text)
        self.assertIn("Measurement only", html_text)
        self.assertNotIn(">Screening<", html_text)
        self.assertIn("Analysis scope", cells)
        self.assertNotIn("Screening", cells)


class ChartDomainTests(unittest.TestCase):
    def test_non_negative_domains_floor_padding_at_zero(self):
        for domain in (DOMAIN_ELAPSED, DOMAIN_INTERVAL, DOMAIN_HOLD, DOMAIN_COUNT):
            self.assertEqual(domain_axis_bounds(-0.7, 10.0, domain=domain, values=[0.3, 9.0]), (0.0, 10.0))
            self.assertEqual(domain_axis_bounds(2.0, 10.0, domain=domain, values=[3.0, 9.0]), (2.0, 10.0))

    def test_negative_data_is_never_cropped(self):
        lo, hi = domain_axis_bounds(-6.0, 10.0, domain=DOMAIN_INTERVAL, values=[-5.0, 9.0])
        self.assertEqual((lo, hi), (-5.0, 10.0))

    def test_signed_domain_untouched(self):
        self.assertEqual(domain_axis_bounds(-3.0, 3.0, domain=DOMAIN_SIGNED, values=[-1.0, 1.0]), (-3.0, 3.0))

    def test_raw_values_not_mutated(self):
        values = [0.1, 0.2, 0.3]
        before = list(values)
        domain_axis_bounds(-1.0, 1.0, domain=DOMAIN_HOLD, values=values)
        self.assertEqual(values, before)

    def _tick_texts(self, svg: str) -> list[str]:
        return re.findall(r"class='tk'[^>]*>([^<]*)<", svg)

    def test_html_interval_series_axis_non_negative(self):
        series = [(float(i * 10), 0.4 + (i % 3) * 0.2) for i in range(20)]
        svg = report_v2._svg_interval_series(series, None)
        ticks = self._tick_texts(svg)
        self.assertTrue(ticks)
        for t in ticks:
            self.assertFalse(t.strip().startswith("-"), ticks)

    def test_html_scatter_axes_non_negative(self):
        points = [(0.5 + i * 0.01, 0.2 + i * 0.01, "") for i in range(20)]
        svg = report_v2._svg_scatter(points, None)
        ticks = self._tick_texts(svg)
        self.assertTrue(ticks)
        for t in ticks:
            self.assertFalse(t.strip().startswith("-"), ticks)

    def test_xlsx_chart_axis_minimums_non_negative(self):
        import zipfile

        events = _events(n=20, interval_ms=0.6, hold_ms=0.3)
        settings = _fixture_settings()
        model = build_completed_result_model(
            locked=settings, events=events, run_id="xl1", target="DUT", operator="op",
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "x.xlsx"
            self.assertTrue(report_v2.write_report_xlsx(path, model))
            with zipfile.ZipFile(path) as zf:
                charts = [zf.read(n).decode("utf-8") for n in zf.namelist() if n.startswith("xl/charts/chart")]
        self.assertTrue(charts)
        mins = [float(v) for xml in charts for v in re.findall(r'<(?:\w+:)?min val="([^"]+)"', xml)]
        self.assertTrue(mins)
        for value in mins:
            self.assertGreaterEqual(value, 0.0, mins)

    def test_html_histogram_count_axis_starts_at_zero(self):
        bins = [{"lo": 0.0, "hi": 1.0, "count": 3}, {"lo": 1.0, "hi": 2.0, "count": 5}]
        svg = report_v2._svg_histogram(bins, None)
        ticks = self._tick_texts(svg)
        self.assertIn("0.0", ticks)
        for t in ticks:
            self.assertFalse(t.strip().startswith("-"), ticks)
        self.assertEqual([b["count"] for b in bins], [3, 5])
