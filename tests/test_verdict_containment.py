"""Phase 4D-2c: the five Custom Limits cannot change the authoritative completed-run status."""
from __future__ import annotations

import json
import os
import random
import re
import sys
import tempfile
import unittest
from pathlib import Path

_SRC = Path(__file__).resolve().parents[1] / "src"
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))
_TESTS = Path(__file__).resolve().parent
if str(_TESTS) not in sys.path:
    sys.path.insert(0, str(_TESTS))

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook

from mouse_click_stability_analyzer import app as legacy_app
from mouse_click_stability_analyzer import report_v2
from mouse_click_stability_analyzer.completed_result import (
    RUN_STATUS_COMPLETE,
    RUN_STATUS_INVALID,
    RUN_STATUS_NO_DATA,
    build_completed_result_model,
    completed_run_status,
    qt_visible_verdict,
    run_status_for_count,
)
from mouse_click_stability_analyzer.host_evidence import build_fixture_host_evidence
from test_qt_limit_semantics import _ARBITRARY_LIMITS, _qt_runtime_settings

_N = 40
_RAW_METRIC_KEYS = (
    "click_count", "interval_avg_ms", "interval_cv_pct", "interval_p99_ms",
    "hold_avg_ms", "hold_cv_pct", "hold_p99_ms", "measured_hz", "measured_interval_ms",
)
_VERDICT_WORD = re.compile(r"\b(PASS|WARN|FAIL)\b")


def _events(intervals: list[float], holds: list[float]) -> list[dict]:
    out, t = [], 0.0
    for i, (step, hold) in enumerate(zip(intervals, holds)):
        t += step if i else 0.0
        out.append({"index": i + 1, "button": "Left", "accepted": True, "event_type": "click",
                    "press_elapsed_ms": t, "release_elapsed_ms": t + hold, "elapsed_ms": t + hold,
                    "press_interval_ms": step if i else None, "hold_ms": hold})
    return out


def _chatter_events():
    intervals = [100.0] * _N
    intervals[20], intervals[21] = 5.0, 95.0
    return _events(intervals, [40.0] * _N)


def _double_click_events():
    intervals = [100.0] * _N
    intervals[20], intervals[21] = 30.0, 70.0
    return _events(intervals, [40.0] * _N)


def _short_actuation_events():
    holds = [40.0] * _N
    holds[20] = 6.0
    return _events([100.0] * _N, holds)


def _low_hold_events():
    holds = [40.0] * _N
    holds[20] = holds[25] = 12.0
    return _events([100.0] * _N, holds)


def _jitter_events():
    rng = random.Random(1)
    return _events([100.0 + rng.uniform(-15.0, 15.0) for _ in range(_N)], [40.0] * _N)


# (limit key, events factory, permissive value, strict value, tag that only the strict value raises, manual modes)
_LIMIT_CASES = (
    ("chatter_threshold_ms", _chatter_events, "3", "8", "CHATTER_HARD", (False, True)),
    ("duplicate_click_threshold_ms", _double_click_events, "20", "40", "HIGH_PROBABILITY_DOUBLE_CLICK", (False, True)),
    ("fixture_actuation_too_short_ms", _short_actuation_events, "4", "8", "FIXTURE_ACTUATION_TOO_SHORT", (False,)),
    ("low_hold_warning_ms", _low_hold_events, "10", "15", "LOW_HOLD_WARNING", (False, True)),
    ("interval_cv_limit_pct", _jitter_events, "30", "5", "INTERVAL_JITTER_WARN", (False, True)),
)


def _settings(manual: bool, **limit_overrides):
    return _qt_runtime_settings(manual=manual, raw_limits=dict(_ARBITRARY_LIMITS, **limit_overrides))


def _model(settings, events, host_evidence_snapshot=None):
    return build_completed_result_model(
        locked=settings, events=events, run_id="contain", target="DUT",
        host_evidence_snapshot=host_evidence_snapshot,
    )


def _raw_measurements(model) -> dict:
    metrics = model["assessment"]["metrics"]
    rows = [
        (row["index"], row["elapsed_ms"], row["button"], row["interval_ms"], row["hold_ms"],
         row["reference_delta_ms"])
        for row in model["click_data"]
    ]
    return {
        "metrics": {key: metrics.get(key) for key in _RAW_METRIC_KEYS},
        "rows": rows,
        "intervals": model["analysis"]["intervals"],
        "series": model["analysis"]["interval_series"],
        "host_press_count": model["run"]["host_press_count"],
    }


class FiveLimitStatusInvarianceTests(unittest.TestCase):
    def test_each_limit_changes_classification_but_never_run_status(self):
        for key, factory, permissive, strict, strict_tag, modes in _LIMIT_CASES:
            for manual in modes:
                with self.subTest(limit=key, manual=manual):
                    events = factory()
                    loose = _model(_settings(manual, **{key: permissive}), factory())
                    tight = _model(_settings(manual, **{key: strict}), factory())
                    self.assertNotIn(strict_tag, loose["assessment"]["tags_all"])
                    self.assertIn(strict_tag, tight["assessment"]["tags_all"])
                    self.assertEqual(_raw_measurements(loose), _raw_measurements(tight))
                    self.assertEqual(loose["run"]["run_status"], RUN_STATUS_COMPLETE)
                    self.assertEqual(tight["run"]["run_status"], RUN_STATUS_COMPLETE)
                    self.assertNotIn("verdict", loose["assessment"])
                    self.assertNotIn("verdict", tight["assessment"])
                    self.assertEqual(factory(), events)

    def test_limit_driven_legacy_verdict_flips_do_not_reach_run_status(self):
        flips = (
            ("chatter_threshold_ms", _chatter_events, "3", "8", True, ("WARN", "FAIL")),
            ("duplicate_click_threshold_ms", _double_click_events, "20", "40", False, ("WARN", "FAIL")),
            ("low_hold_warning_ms", _low_hold_events, "10", "15", False, ("PASS", "WARN")),
        )
        for key, factory, permissive, strict, manual, expected_legacy in flips:
            with self.subTest(limit=key, manual=manual):
                legacy, statuses = [], []
                for value in (permissive, strict):
                    settings = _settings(manual, **{key: value})
                    model = _model(settings, factory())
                    a = model["assessment"]
                    legacy.append(qt_visible_verdict(a["tags_all"], a["metrics"], settings))
                    statuses.append(model["run"]["run_status"])
                self.assertEqual(tuple(legacy), expected_legacy)
                self.assertEqual(statuses, [RUN_STATUS_COMPLETE, RUN_STATUS_COMPLETE])

    def test_all_five_limits_strict_vs_permissive_share_status_and_measurements(self):
        events = _chatter_events()
        permissive = {key: p for key, _f, p, _s, _t, _m in _LIMIT_CASES}
        strict = {key: s for key, _f, _p, s, _t, _m in _LIMIT_CASES}
        for manual in (False, True):
            with self.subTest(manual=manual):
                loose = _model(_settings(manual, **permissive), list(events))
                tight = _model(_settings(manual, **strict), list(events))
                self.assertNotEqual(loose["assessment"]["tags_all"], tight["assessment"]["tags_all"])
                self.assertEqual(_raw_measurements(loose), _raw_measurements(tight))
                self.assertEqual(loose["run"]["run_status"], tight["run"]["run_status"])
                self.assertEqual(loose["run"]["run_status"], RUN_STATUS_COMPLETE)

    def test_completed_run_status_signature_has_no_settings_input(self):
        events = _chatter_events()
        self.assertEqual(completed_run_status(events, ["CHATTER_HARD", "INTERVAL_JITTER_WARN"]), RUN_STATUS_COMPLETE)
        self.assertEqual(completed_run_status(events, []), RUN_STATUS_COMPLETE)


class FixtureTimingToleranceTests(unittest.TestCase):
    def test_tolerance_changes_host_evidence_but_not_run_status(self):
        events = _jitter_events()
        presses = [e["press_elapsed_ms"] for e in events]
        duration_s = (presses[-1] + 50.0) / 1000.0
        snapshots = [build_fixture_host_evidence(duration_s, 10.0, tol, presses) for tol in (1.0, 20.0)]
        self.assertNotEqual(snapshots[0], snapshots[1])
        models = [
            _model(_settings(False) | {"fixture_timing_tolerance_ms": tol}, list(events), snap)
            for tol, snap in zip((1.0, 20.0), snapshots)
        ]
        self.assertEqual(models[0]["host_evidence"]["snapshot"], snapshots[0])
        self.assertEqual(models[1]["host_evidence"]["snapshot"], snapshots[1])
        self.assertEqual(_raw_measurements(models[0]), _raw_measurements(models[1]))
        self.assertEqual([m["run"]["run_status"] for m in models], [RUN_STATUS_COMPLETE] * 2)


class NoDataAndInvalidStatusTests(unittest.TestCase):
    def test_zero_accepted_presses_is_no_data_not_fail(self):
        self.assertEqual(run_status_for_count(0), RUN_STATUS_NO_DATA)
        self.assertEqual(run_status_for_count(None, ["CHATTER_HARD"]), RUN_STATUS_NO_DATA)
        self.assertEqual(completed_run_status([]), RUN_STATUS_NO_DATA)
        rejected = [dict(e, accepted=False) for e in _chatter_events()]
        self.assertEqual(completed_run_status(rejected), RUN_STATUS_NO_DATA)

    def test_capture_validity_tags_mark_run_invalid(self):
        for tag in ("CAPTURE_SOURCE_INVALID", "WRONG_BUTTON_DETECTED"):
            with self.subTest(tag=tag):
                self.assertEqual(run_status_for_count(10, [tag]), RUN_STATUS_INVALID)

    def test_heuristic_tags_never_produce_invalid_or_no_data(self):
        heuristic = [
            "CHATTER_HARD", "HIGH_PROBABILITY_DOUBLE_CLICK", "TRUE_EXTRA_CLICK", "FIXTURE_ACTUATION_TOO_SHORT",
            "INCOMPLETE_REBOUND", "LOW_HOLD_WARNING", "INTERVAL_JITTER_WARN", "TIMING_JITTER_PAIR",
            "MISS_OR_SLOW", "LONG_INTERVAL_SPIKE_WARN", "INSUFFICIENT_CLICKS",
        ]
        self.assertEqual(run_status_for_count(10, heuristic), RUN_STATUS_COMPLETE)


class ReportNeutralStatusTests(unittest.TestCase):
    def _visible(self, html: str) -> str:
        html = re.sub(r"<style>.*?</style>", " ", html, flags=re.S)
        return re.sub(r"<[^>]+>", " ", html)

    def test_report_v2_renders_complete_for_strict_and_permissive_limits(self):
        for value in ("3", "8"):
            with self.subTest(chatter=value):
                model = _model(_settings(False, chatter_threshold_ms=value), _chatter_events())
                with tempfile.TemporaryDirectory() as tmp:
                    html_path, xlsx_path = Path(tmp) / "r.html", Path(tmp) / "r.xlsx"
                    report_v2.write_report_html(html_path, model)
                    self.assertTrue(report_v2.write_report_xlsx(xlsx_path, model))
                    visible = self._visible(html_path.read_text(encoding="utf-8"))
                    self.assertIn("Run Status", visible)
                    self.assertIn("COMPLETE", visible)
                    self.assertNotIn("Test Result", visible)
                    self.assertIsNone(_VERDICT_WORD.search(visible))
                    wb = load_workbook(xlsx_path)
                    try:
                        cells = [str(c) for ws in wb for row in ws.iter_rows(values_only=True) for c in row if c]
                    finally:
                        wb.close()
                    self.assertIn("COMPLETE", cells)
                    self.assertNotIn("Test Result", cells)
                    self.assertFalse([c for c in cells if _VERDICT_WORD.search(c)])

    def test_json_sidecars_separate_legacy_verdict_from_run_status(self):
        summaries = []
        for value in ("3", "8"):
            settings = _settings(True, chatter_threshold_ms=value)
            events = _chatter_events()
            model = _model(settings, events)
            with tempfile.TemporaryDirectory() as tmp:
                paths = legacy_app.write_reports(Path(tmp), "DUT", "op", settings, events, f"run_{value}",
                                                 completed_model=model)
                summary = json.loads(Path(paths["summary_json"]).read_text(encoding="utf-8"))
                detail = json.loads(Path(paths["detail_json"]).read_text(encoding="utf-8"))
                text = Path(paths["txt"]).read_text(encoding="utf-8")
            self.assertEqual(summary["run_status"], RUN_STATUS_COMPLETE)
            self.assertEqual(summary["result"], RUN_STATUS_COMPLETE)
            self.assertEqual(summary["result_semantics"], "run_status")
            self.assertEqual(detail["run_status"], RUN_STATUS_COMPLETE)
            self.assertIs(summary["legacy_heuristic_assessment"]["authoritative"], False)
            self.assertIs(detail["legacy_heuristic_assessment"]["authoritative"], False)
            self.assertNotIn("health", summary["metrics"])
            self.assertIn("Run Status: COMPLETE", text)
            summaries.append(summary)
        self.assertEqual(
            [s["legacy_heuristic_assessment"]["result"] for s in summaries], ["WARN", "FAIL"],
        )
        self.assertEqual([s["run_status"] for s in summaries], [RUN_STATUS_COMPLETE] * 2)


def _make_qapp():
    try:
        from PySide6.QtWidgets import QApplication
    except Exception:  # pragma: no cover
        return None
    return QApplication.instance() or QApplication([])


class QtRunStatusContainmentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._qapp = _make_qapp()
        if cls._qapp is None:
            raise unittest.SkipTest("PySide6 QApplication unavailable")
        from mouse_click_stability_analyzer.qt_shell import app as qt_app
        cls.qt_app = qt_app
        cls._tmp = tempfile.TemporaryDirectory()

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def _window(self):
        from PySide6.QtCore import QCoreApplication, QEvent
        win = self.qt_app.MainWindow(report_root=Path(self._tmp.name))
        win._listen_start_timer.stop()
        win.ui_timer.stop()
        win.ensure_raw_input = lambda: None
        win._cursor_in_test_area = lambda: True
        win._cursor_over_capture_exclusion_ui = lambda: False
        win.listening = True

        def cleanup():
            win.close()
            win.deleteLater()
            self._qapp.processEvents()
            QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)

        self.addCleanup(cleanup)
        return win

    def _run(self, *, manual: bool, limits: dict[str, str], offsets_ms, tol: str = "2"):
        win = self._window()
        if manual:
            win.radio_manual.setChecked(True)
        else:
            win.radio_fixture.setChecked(True)
            win.hz_combo.setCurrentText("40")
            win.timing_tol_edit.setText(tol)
            win.auto_stop_on.setChecked(True)
            win.auto_stop_sec_combo.setCurrentText("1")
        for key, value in limits.items():
            raise AssertionError(f"Setup no longer exposes custom-limit editors for {key}={value}")
        self.assertTrue(win.lock_profile())
        self.assertTrue(win.arm_capture())
        t0 = float(win._arm_boundary_perf) + 0.001
        win._begin_session_from_first_click("Left", t0, t0 + 0.01, {})
        self.assertTrue(win.running)
        for off in offsets_ms[1:]:
            press = t0 + off / 1000.0
            win.record_click("Left", press, press + 0.01, {})
        win.stop_run()
        if win._post_run_input_guard_active:
            win._release_post_run_input_guard()
        return win

    def test_setup_cannot_vary_legacy_heuristic_tags(self):
        offsets = [0, 25, 30, 50, 75, 100, 125, 150]
        defaults = self.qt_app.qt_custom_limit_defaults()
        for manual in (False, True):
            with self.subTest(manual=manual):
                a = self._run(manual=manual, limits={}, offsets_ms=offsets)
                b = self._run(manual=manual, limits={}, offsets_ms=offsets)
                fa, fb = a._completed_run_facts(), b._completed_run_facts()
                self.assertEqual(fa["model"]["assessment"]["tags_all"], fb["model"]["assessment"]["tags_all"])
                self.assertEqual(_raw_measurements(fa["model"]), _raw_measurements(fb["model"]))
                for win, facts in ((a, fa), (b, fb)):
                    self.assertEqual(facts["run_status"], RUN_STATUS_COMPLETE)
                    self.assertEqual(win.val_verdict.text(), RUN_STATUS_COMPLETE)
                    for key, value in defaults.items():
                        self.assertEqual(win.profile_locked_settings[key], value)
                    win._refresh_results_workspace()
                    self.assertEqual(win.results_verdict_label.text(), RUN_STATUS_COMPLETE)

    def test_timing_tolerance_changes_host_evidence_not_run_status(self):
        offsets = [0, 27, 49, 77, 98, 127, 150, 173]
        a = self._run(manual=False, limits={}, offsets_ms=offsets, tol="1")
        b = self._run(manual=False, limits={}, offsets_ms=offsets, tol="10")
        self.assertIsNotNone(a.host_evidence_snapshot)
        self.assertIsNotNone(b.host_evidence_snapshot)
        self.assertNotEqual(a.host_evidence_snapshot, b.host_evidence_snapshot)
        for win in (a, b):
            self.assertEqual(win._completed_run_facts()["run_status"], RUN_STATUS_COMPLETE)
            self.assertEqual(win.val_verdict.text(), RUN_STATUS_COMPLETE)


if __name__ == "__main__":
    unittest.main()
