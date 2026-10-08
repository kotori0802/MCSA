"""Phase 4E-B1: machine-readable Schedule Association Window provenance."""
from __future__ import annotations

import json
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path

_SRC = Path(__file__).resolve().parents[1] / "src"
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook

from mouse_click_stability_analyzer import app as legacy_app
from mouse_click_stability_analyzer.completed_result import (
    RUN_STATUS_COMPLETE,
    build_completed_result_model,
)
from mouse_click_stability_analyzer.fixture_matching import match_fixture_press_observations
from mouse_click_stability_analyzer.host_evidence import build_fixture_host_evidence

_WINDOW_KEY = "fixture_timing_tolerance_ms"
_PILOT_ID = "20261005_172910"
_REPO = Path(__file__).resolve().parents[1]


def _events(n: int = 20, interval_ms: float = 100.0, hold_ms: float = 40.0) -> list[dict]:
    """Same time origins as ``record_click``: press from first press, release from first release."""
    presses = [i * interval_ms for i in range(n)]
    releases = [p + hold_ms for p in presses]
    out = []
    for i, (press, release) in enumerate(zip(presses, releases)):
        release_elapsed = release - releases[0]
        out.append({
            "index": i + 1, "button": "Left", "accepted": True, "event_type": "click",
            "press_elapsed_ms": press - presses[0],
            "release_elapsed_ms": release_elapsed, "elapsed_ms": release_elapsed,
            "press_interval_ms": None if i == 0 else press - presses[i - 1],
            "hold_ms": release - press,
        })
    return out


def _fixture_settings(window_ms: float = 12.49, **extra) -> dict:
    settings = {
        "profile_name": "Fixture_Custom",
        "manual_mode": False,
        "focus_button": "Left",
        "button": "Left",
        "fixture_hz": 10.0,
        "fixture_reference_hz": 10.0,
        "duration_sec": 2.0,
        "test_profile": "Fixture Custom",
        "capture_mode": "Windows Raw Input / 裝置RawInput",
        _WINDOW_KEY: window_ms,
    }
    settings.update(extra)
    return settings


def _manual_settings() -> dict:
    return {
        "profile_name": "Manual_Left_Stable",
        "manual_mode": True,
        "focus_button": "Left",
        "button": "Left",
        "fixture_hz": 0.0,
        "fixture_reference_hz": 0.0,
        "duration_sec": 2.0,
        "test_profile": "Manual Stable",
        "capture_mode": "Windows Raw Input / 裝置RawInput",
        _WINDOW_KEY: None,
    }


def _html_info(html_text: str) -> dict[str, str]:
    match = re.search(r"<section id=['\"]test-information['\"]>(.*?)</section>", html_text, flags=re.S)
    if not match:
        return {}
    pairs = re.findall(r"<th scope='row'>(.*?)</th><td[^>]*>(.*?)</td>", match.group(1), flags=re.S)
    strip = lambda s: re.sub(r"<[^>]+>", "", s)
    return {strip(k): strip(v) for k, v in pairs}


def _xlsx_summary(path: Path) -> dict[str, object]:
    wb = load_workbook(path)
    try:
        out: dict[str, object] = {}
        for row in wb["Summary"].iter_rows(min_col=1, max_col=2, values_only=True):
            if row[0]:
                out[str(row[0])] = row[1]
        return out
    finally:
        wb.close()


def _write(settings, events, run_id, *, model=None):
    with tempfile.TemporaryDirectory() as tmp:
        if model is None:
            model = build_completed_result_model(
                locked=settings, events=events, run_id=run_id, target="DUT", operator="op",
            )
        paths = legacy_app.write_reports(
            Path(tmp), "DUT", "op", settings, events, run_id, completed_model=model,
        )
        evidence = json.loads(Path(paths["evidence_json"]).read_text(encoding="utf-8"))
        html_text = Path(paths["engineering_html"]).read_text(encoding="utf-8")
        xlsx_kv = _xlsx_summary(Path(paths["engineering_xlsx"])) if paths.get("engineering_xlsx") and Path(paths["engineering_xlsx"]).is_file() else {}
        return evidence, evidence, html_text, xlsx_kv, model


def _model_run(value, *, frozen_settings):
    return {"run": {"timing_tolerance_ms": value, "settings": frozen_settings}}


class EvidenceWhitelistTests(unittest.TestCase):
    def test_standard_evidence_carries_4eb_fields(self):
        evidence, _, *_ = _write(_fixture_settings(12.49), _events(), "char_4eb")
        self.assertEqual(evidence["schema"], "MCSA_EVIDENCE_V1")
        self.assertEqual(evidence["run_status"], RUN_STATUS_COMPLETE)
        self.assertTrue(evidence["tool_version"])
        settings = evidence["settings"]
        for key in (
            "manual_mode", "focus_button", "duration_sec",
            "fixture_reference_hz", "fixture_timing_tolerance_ms",
        ):
            self.assertIn(key, settings, key)
        self.assertFalse(settings["manual_mode"])
        self.assertEqual(settings["focus_button"], "Left")
        self.assertEqual(settings["duration_sec"], 2.0)
        self.assertEqual(settings["fixture_reference_hz"], 10.0)
        self.assertEqual(settings["fixture_timing_tolerance_ms"], 12.49)
        events = evidence["events"]
        self.assertGreaterEqual(len(events), 1)
        for event in events:
            self.assertIn("accepted", event)
            self.assertIn("button", event)
            self.assertIn("press_elapsed_ms", event)
            self.assertNotIn("click_class", event)
            self.assertNotIn("issue_tags", event)
            self.assertNotIn("expected_click_index", event)
            self.assertNotIn("drift_ms", event)


class SyntheticFixtureTimeOriginTests(unittest.TestCase):
    def test_events_follow_record_click_time_origins(self):
        hold = 40.0
        events = _events(5, 100.0, hold)
        self.assertEqual([e["press_elapsed_ms"] for e in events], [0.0, 100.0, 200.0, 300.0, 400.0])
        # Constant hold: release relative to the first release tracks press relative to the first press.
        self.assertEqual([e["release_elapsed_ms"] for e in events], [0.0, 100.0, 200.0, 300.0, 400.0])
        self.assertEqual({e["hold_ms"] for e in events}, {hold})
        for e in events:
            self.assertEqual(
                e["release_elapsed_ms"], e["press_elapsed_ms"] + e["hold_ms"] - events[0]["hold_ms"],
            )


class FrozenWindowHelperTests(unittest.TestCase):
    def test_uses_frozen_model_value_and_provenance_only(self):
        stamp = legacy_app.stamp_frozen_schedule_association_window
        self.assertEqual(stamp({}, _model_run(12.49, frozen_settings={_WINDOW_KEY: 12.49}))[_WINDOW_KEY], 12.49)
        self.assertEqual(stamp({}, _model_run(0.0, frozen_settings={_WINDOW_KEY: 0.0}))[_WINDOW_KEY], 0.0)
        self.assertIs(stamp({}, _model_run(None, frozen_settings={_WINDOW_KEY: None}))[_WINDOW_KEY], None)
        self.assertNotIn(_WINDOW_KEY, stamp({}, _model_run(None, frozen_settings={})))
        self.assertNotIn(
            _WINDOW_KEY, stamp({}, _model_run(None, frozen_settings={"chatter_threshold_ms": 8.0})),
        )
        # A stale value already on the normalized settings is never kept.
        self.assertNotIn(_WINDOW_KEY, stamp({_WINDOW_KEY: 5.0}, _model_run(None, frozen_settings={})))
        self.assertEqual(
            stamp({_WINDOW_KEY: 5.0}, _model_run(7.25, frozen_settings={_WINDOW_KEY: 5.0}))[_WINDOW_KEY], 7.25,
        )


class FixturePersistenceTests(unittest.TestCase):
    def test_window_is_identical_in_detail_and_summary_json(self):
        summary, detail, html_text, xlsx, model = _write(_fixture_settings(12.49), _events(), "win_1249")
        self.assertIn(_WINDOW_KEY, summary["settings"])
        self.assertIn(_WINDOW_KEY, detail["settings"])
        self.assertEqual(summary["settings"][_WINDOW_KEY], 12.49)
        self.assertEqual(detail["settings"][_WINDOW_KEY], 12.49)
        self.assertEqual(summary["settings"][_WINDOW_KEY], detail["settings"][_WINDOW_KEY])
        self.assertEqual(model["run"]["timing_tolerance_ms"], 12.49)
        self.assertEqual(_html_info(html_text)["Schedule Association Window"], "±12.49 ms")
        self.assertEqual(xlsx["Schedule Association Window"], "±12.49 ms")

    def test_precision_is_not_truncated_or_recovered_from_report_text(self):
        # 7.251 displays as ±7.25 ms (2 dp). JSON must keep 7.251.
        summary, detail, html_text, xlsx, model = _write(_fixture_settings(7.251), _events(), "win_7251")
        stored = summary["settings"][_WINDOW_KEY]
        self.assertEqual(stored, 7.251)
        self.assertEqual(detail["settings"][_WINDOW_KEY], 7.251)
        self.assertNotEqual(stored, 7.25)
        self.assertNotEqual(stored, 7)
        self.assertEqual(_html_info(html_text)["Schedule Association Window"], "±7.25 ms")
        self.assertEqual(model["run"]["timing_tolerance_ms"], stored)
        self.assertIn('"fixture_timing_tolerance_ms": 7.251', json.dumps(summary["settings"]))

    def test_divergent_caller_settings_cannot_split_the_bundle(self):
        settings = _fixture_settings(5.0)
        events = _events()
        model = build_completed_result_model(
            locked=dict(settings), events=events, run_id="diverge", target="DUT", operator="op",
        )
        model["run"]["timing_tolerance_ms"] = 7.25
        summary, detail, html_text, xlsx, _ = _write(settings, events, "diverge", model=model)
        self.assertEqual(settings[_WINDOW_KEY], 5.0)
        self.assertEqual(summary["settings"][_WINDOW_KEY], 7.25)
        self.assertEqual(detail["settings"][_WINDOW_KEY], 7.25)
        self.assertEqual(_html_info(html_text)["Schedule Association Window"], "±7.25 ms")
        self.assertEqual(xlsx["Schedule Association Window"], "±7.25 ms")

    def test_caller_window_without_frozen_provenance_is_not_persisted(self):
        frozen = _fixture_settings()
        del frozen[_WINDOW_KEY]
        events = _events()
        model = build_completed_result_model(
            locked=frozen, events=events, run_id="no_prov", target="DUT", operator="op",
        )
        summary, detail, html_text, *_ = _write(_fixture_settings(5.0), events, "no_prov", model=model)
        self.assertNotIn(_WINDOW_KEY, summary["settings"])
        self.assertNotIn(_WINDOW_KEY, detail["settings"])
        self.assertNotIn("Schedule Association Window", _html_info(html_text))

    def test_fallback_model_path_still_persists_the_window(self):
        events = _events()
        with tempfile.TemporaryDirectory() as tmp:
            paths = legacy_app.write_reports(Path(tmp), "DUT", "op", _fixture_settings(7.251), events, "fallback")
            evidence = json.loads(Path(paths["evidence_json"]).read_text(encoding="utf-8"))
            html_text = Path(paths["engineering_html"]).read_text(encoding="utf-8")
        self.assertEqual(evidence["settings"][_WINDOW_KEY], 7.251)
        self.assertEqual(_html_info(html_text)["Schedule Association Window"], "±7.25 ms")

    def test_zero_is_a_real_configured_window(self):
        summary, detail, *_ = _write(_fixture_settings(0.0), _events(), "win_zero")
        self.assertEqual(summary["settings"][_WINDOW_KEY], 0.0)
        self.assertEqual(detail["settings"][_WINDOW_KEY], 0.0)


class ManualModeTests(unittest.TestCase):
    def test_manual_persists_null_not_zero(self):
        summary, detail, html_text, xlsx, model = _write(_manual_settings(), _events(), "win_manual")
        self.assertIn(_WINDOW_KEY, summary["settings"])
        self.assertIn(_WINDOW_KEY, detail["settings"])
        self.assertIsNone(summary["settings"][_WINDOW_KEY])
        self.assertIsNone(detail["settings"][_WINDOW_KEY])
        self.assertIsNone(model["run"]["timing_tolerance_ms"])
        self.assertNotIn("Schedule Association Window", _html_info(html_text))
        self.assertNotIn("Schedule Association Window", xlsx)


class BackwardCompatibilityTests(unittest.TestCase):
    def test_missing_key_stays_missing_and_is_not_a_fallback(self):
        settings = _fixture_settings()
        del settings[_WINDOW_KEY]
        settings["chatter_threshold_ms"] = 8.0
        settings["duplicate_click_threshold_ms"] = 30.0
        summary, detail, *_ = _write(settings, _events(), "win_old")
        self.assertNotIn(_WINDOW_KEY, summary["settings"])
        self.assertNotIn(_WINDOW_KEY, detail["settings"])
        self.assertNotIn("chatter_threshold_ms", summary["settings"])
        self.assertNotIn("duplicate_click_threshold_ms", summary["settings"])

    def test_pilot_artifact_lacks_native_json_window(self):
        detail_path = _REPO / "reports" / f"mcsa_{_PILOT_ID}_detail.json"
        summary_path = _REPO / "reports" / f"mcsa_{_PILOT_ID}_summary.json"
        html_path = _REPO / "reports" / f"mcsa_{_PILOT_ID}_report.html"
        if not detail_path.is_file() or not summary_path.is_file():
            self.skipTest(f"pilot artifacts for {_PILOT_ID} are absent")
        detail = json.loads(detail_path.read_text(encoding="utf-8"))
        summary = json.loads(summary_path.read_text(encoding="utf-8"))
        self.assertNotIn(_WINDOW_KEY, detail.get("settings") or {})
        self.assertNotIn(_WINDOW_KEY, summary.get("settings") or {})
        self.assertTrue(html_path.is_file())
        html_text = html_path.read_text(encoding="utf-8")
        self.assertEqual(_html_info(html_text).get("Schedule Association Window"), "±20.00 ms")
        xlsx_path = _REPO / "reports" / f"mcsa_{_PILOT_ID}_report.xlsx"
        if xlsx_path.is_file():
            self.assertEqual(_xlsx_summary(xlsx_path).get("Schedule Association Window"), "±20.00 ms")


_AUTO_STOP_KEY = "auto_stop_enabled"


def _write_bundle(settings, events, run_id, *, model=None):
    with tempfile.TemporaryDirectory() as tmp:
        if model is None:
            model = build_completed_result_model(
                locked=settings, events=events, run_id=run_id, target="DUT", operator="op",
            )
        paths = legacy_app.write_reports(
            Path(tmp), "DUT", "op", settings, events, run_id, completed_model=model,
        )
        package = Path(paths["evidence_json"]).parent
        return {
            "evidence": json.loads(Path(paths["evidence_json"]).read_text(encoding="utf-8")),
            "events_csv": Path(paths["all_click_log_csv"]).read_text(encoding="utf-8"),
            "html_info": _html_info(Path(paths["engineering_html"]).read_text(encoding="utf-8")),
            "xlsx": _xlsx_summary(Path(paths["engineering_xlsx"])) if paths.get("engineering_xlsx") else {},
            "files": sorted(p.name for p in package.iterdir()),
            "model": model,
        }


def _without_generated(info: dict) -> dict:
    return {k: v for k, v in info.items() if "generated" not in str(k).lower()}


class AutoStopEvidenceTests(unittest.TestCase):
    def test_auto_stop_on_serializes_true_with_separate_duration(self):
        out = _write_bundle(_fixture_settings(duration_sec=60.0, **{_AUTO_STOP_KEY: True}), _events(), "as_on")
        settings = out["evidence"]["settings"]
        self.assertIs(settings[_AUTO_STOP_KEY], True)
        self.assertEqual(settings["duration_sec"], 60.0)
        self.assertIs(out["model"]["run"][_AUTO_STOP_KEY], True)
        self.assertIn('"auto_stop_enabled": true', json.dumps(settings))

    def test_auto_stop_off_serializes_false(self):
        out = _write_bundle(_fixture_settings(duration_sec=0, **{_AUTO_STOP_KEY: False}), _events(), "as_off")
        self.assertIs(out["evidence"]["settings"][_AUTO_STOP_KEY], False)
        self.assertIs(out["model"]["run"][_AUTO_STOP_KEY], False)

    def test_boolean_is_independent_of_duration_sec(self):
        out = _write_bundle(_fixture_settings(duration_sec=60.0, **{_AUTO_STOP_KEY: False}), _events(), "as_indep")
        self.assertIs(out["evidence"]["settings"][_AUTO_STOP_KEY], False)
        self.assertEqual(out["evidence"]["settings"]["duration_sec"], 60.0)

    def test_completed_value_is_frozen_against_live_setup_change(self):
        locked = _fixture_settings(duration_sec=60.0, **{_AUTO_STOP_KEY: True})
        events = _events()
        model = build_completed_result_model(
            locked=dict(locked), events=events, run_id="as_frozen", target="DUT", operator="op",
        )
        live = dict(locked, **{_AUTO_STOP_KEY: False})
        out = _write_bundle(live, events, "as_frozen", model=model)
        self.assertIs(out["evidence"]["settings"][_AUTO_STOP_KEY], True)

    def test_old_input_without_field_stays_absent(self):
        out = _write_bundle(_fixture_settings(duration_sec=60.0), _events(), "as_old")
        self.assertNotIn(_AUTO_STOP_KEY, out["evidence"]["settings"])
        self.assertIsNone(out["model"]["run"][_AUTO_STOP_KEY])

    def test_non_boolean_frozen_value_is_not_persisted(self):
        out = _write_bundle(_fixture_settings(**{_AUTO_STOP_KEY: "yes"}), _events(), "as_bad")
        self.assertNotIn(_AUTO_STOP_KEY, out["evidence"]["settings"])

    def test_stamp_reads_frozen_model_only(self):
        stamp = legacy_app.stamp_frozen_auto_stop_enabled
        self.assertIs(stamp({}, {"run": {_AUTO_STOP_KEY: True}})[_AUTO_STOP_KEY], True)
        self.assertIs(stamp({}, {"run": {_AUTO_STOP_KEY: False}})[_AUTO_STOP_KEY], False)
        self.assertNotIn(_AUTO_STOP_KEY, stamp({}, {"run": {_AUTO_STOP_KEY: None}}))
        self.assertNotIn(_AUTO_STOP_KEY, stamp({_AUTO_STOP_KEY: True}, {"run": {}}))
        self.assertNotIn(_AUTO_STOP_KEY, stamp({"duration_sec": 60.0}, {"run": {"auto_stop_sec": 60.0}}))

    def test_field_does_not_change_other_artifacts(self):
        events = _events()
        base = _fixture_settings(duration_sec=60.0)
        with_field = _write_bundle(dict(base, **{_AUTO_STOP_KEY: True}), events, "as_same")
        without = _write_bundle(dict(base), events, "as_same")
        self.assertEqual(with_field["events_csv"], without["events_csv"])
        self.assertEqual(with_field["files"], without["files"])
        self.assertEqual(len(with_field["files"]), 4)
        self.assertEqual(_without_generated(with_field["html_info"]), _without_generated(without["html_info"]))
        self.assertEqual(_without_generated(with_field["xlsx"]), _without_generated(without["xlsx"]))
        ev_a, ev_b = with_field["evidence"], without["evidence"]
        self.assertEqual(ev_a["schema"], "MCSA_EVIDENCE_V1")
        self.assertEqual(ev_a["run_status"], ev_b["run_status"])
        self.assertEqual(ev_a["measurements"], ev_b["measurements"])
        self.assertEqual(ev_a["events"], ev_b["events"])
        self.assertEqual(ev_a.get("interval_comparison"), ev_b.get("interval_comparison"))
        self.assertEqual(ev_a.get("host_evidence"), ev_b.get("host_evidence"))
        a_settings = dict(ev_a["settings"])
        a_settings.pop(_AUTO_STOP_KEY)
        self.assertEqual(a_settings, ev_b["settings"])


class LegacyAndMatcherContainmentTests(unittest.TestCase):
    def test_authoritative_run_status_and_no_legacy_verdict_block(self):
        settings = _fixture_settings(12.49)
        events = _events()
        summary, detail, *_rest = _write(settings, events, "win_legacy")
        self.assertEqual(summary["run_status"], RUN_STATUS_COMPLETE)
        self.assertEqual(detail["run_status"], RUN_STATUS_COMPLETE)
        self.assertNotIn("result", summary)
        self.assertNotIn("legacy_heuristic_assessment", summary)
        self.assertNotIn("legacy_heuristic_assessment", detail)
        self.assertIn(summary["run_status"], (RUN_STATUS_COMPLETE, "NO_DATA", "INVALID"))
        self.assertNotIn(summary["run_status"], ("PASS", "WARN", "FAIL"))

    def test_adding_the_field_does_not_change_factual_measurements(self):
        events = _events()
        base = _fixture_settings(12.49)
        without = dict(base)
        del without[_WINDOW_KEY]
        s_with, _, *_ = _write(base, events, "with_w")
        s_without, _, *_ = _write(without, events, "without_w")
        for key in ("host_press_count", "interval_avg_ms", "interval_count"):
            self.assertEqual(s_with["measurements"].get(key), s_without["measurements"].get(key), key)
        self.assertNotIn("legacy_heuristic_assessment", s_with)
        self.assertNotIn("legacy_heuristic_assessment", s_without)

    def test_matcher_and_host_evidence_arithmetic_unchanged(self):
        times = [e["press_elapsed_ms"] for e in _events()]
        matching = match_fixture_press_observations(2.0, 10.0, 12.49, times)
        evidence = build_fixture_host_evidence(2.0, 10.0, 12.49, times)
        self.assertEqual(matching["timing_tolerance_ms"], 12.49)
        self.assertEqual(matching["matched_observation_count"], 20)
        self.assertEqual(matching["unmatched_observation_count"], 0)
        self.assertEqual(matching["ambiguous_observation_count"], 0)
        self.assertEqual(evidence["matching"]["matched_host_observation_count"], 20)
        self.assertEqual(evidence["configuration"]["timing_tolerance_ms"], 12.49)


if __name__ == "__main__":
    unittest.main()
