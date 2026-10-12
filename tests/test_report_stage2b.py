"""Stage 2B: report elapsed, histogram N, Measurement-only columns."""
from __future__ import annotations

import inspect
import math
import os
import sys
import tempfile
import unittest
from pathlib import Path

_SRC = Path(__file__).resolve().parents[1] / "src"
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from datetime import timedelta

from openpyxl import load_workbook

from mouse_click_stability_analyzer.completed_result import (
    build_completed_result_model,
    count_histogram,
    histogram_bin_edges,
    interval_histogram_bins,
)
from mouse_click_stability_analyzer.duration import EXCEL_ELAPSED_NUMBER_FORMAT, excel_elapsed_days, format_elapsed_hms
from mouse_click_stability_analyzer.host_inventory import collect_host_information
from mouse_click_stability_analyzer.report_v2 import (
    click_data_columns,
    render_report_html,
    write_report_html,
    write_report_xlsx,
)
from mouse_click_stability_analyzer import report_v2


def _settings(*, manual=False):
    if manual:
        return {"manual_mode": True, "focus_button": "Left", "profile_name": "Manual", "fixture_reference_hz": 0}
    return {
        "manual_mode": False, "focus_button": "Left", "profile_name": "Fixture",
        "fixture_reference_hz": 10.0, "fixture_hz": 10.0, "fixture_timing_tolerance_ms": 2.0,
        "duration_sec": 2.0,
    }


def _events(intervals):
    rows = []
    t = 0.0
    for i, step in enumerate([0.0] + list(intervals)):
        t += step if i else 0.0
        rows.append({
            "index": i + 1, "click_index": i + 1, "button": "Left", "accepted": True,
            "event_type": "click", "press_elapsed_ms": t, "release_elapsed_ms": t + 8.0,
            "elapsed_ms": t + 8.0, "press_interval_ms": None if i == 0 else step, "hold_ms": 8.0,
        })
    return rows


class HistogramContractTests(unittest.TestCase):
    def test_n_is_included_finite_intervals_and_sum_matches(self):
        events = _events([10.0, 12.0, 11.0, 10.5])
        model = build_completed_result_model(
            locked=_settings(manual=True), events=events, run_id="hist", target="T",
        )
        included = model["analysis"]["intervals"]
        bins = model["analysis"]["histogram"]
        self.assertEqual(len(included), 4)
        self.assertEqual(sum(int(b["count"]) for b in bins), 4)

    def test_empty_and_single_press(self):
        self.assertEqual(interval_histogram_bins([]), [])
        model = build_completed_result_model(
            locked=_settings(manual=True), events=_events([]), run_id="one", target="T",
        )
        self.assertEqual(model["analysis"]["intervals"], [])
        self.assertEqual(model["analysis"]["histogram"], [])
        self.assertEqual(len(model["click_data"]), 1)

    def test_algorithm_unchanged_edges_and_last_inclusive(self):
        values = [10.0, 20.0, 30.0, 40.0, 50.0, 60.0]
        edges = histogram_bin_edges(values)
        self.assertEqual(edges[0][0], 10.0)
        self.assertEqual(edges[-1][1], 60.0)
        self.assertIn("last bin includes hi", inspect.getdoc(count_histogram) or "")
        bins = count_histogram(values, edges)
        self.assertEqual(sum(int(b["count"]) for b in bins), 6)

    def test_model_n_matches_press_interval_population(self):
        events = _events([100.0, 101.0, 99.0])
        model = build_completed_result_model(locked=_settings(manual=True), events=events, run_id="n", target="T")
        n = len(model["analysis"]["intervals"])
        self.assertEqual(n, 3)
        self.assertEqual(sum(int(b["count"]) for b in model["analysis"]["histogram"]), n)
        view = report_v2.build_report_view(model)
        self.assertEqual(view["interval_observation_count"], 3)
        html = render_report_html(model, view)
        self.assertIn("Click interval distribution (N=3)", html)
        self.assertIn("N is the number of interval observations included", html)


class MeasurementOnlyColumnsTests(unittest.TestCase):
    def test_xlsx_omits_observation_columns(self):
        cols = click_data_columns(False)
        self.assertNotIn("Observation", cols)
        self.assertNotIn("Observation flag(s)", cols)
        self.assertIn("Elapsed time", cols)
        self.assertNotIn("Elapsed (ms)", cols)
        self.assertIn("Interval Δ vs ref. (ms)", cols)
        manual_cols = click_data_columns(True)
        self.assertNotIn("Interval Δ vs ref. (ms)", manual_cols)
        future = click_data_columns(False, include_classification=True)
        self.assertIn("Observation", future)
        self.assertIn("Observation flag(s)", future)

    def test_xlsx_and_html_measurement_only(self):
        model = build_completed_result_model(
            locked=_settings(manual=True), events=_events([100.0, 120.0]), run_id="cols", target="T",
        )
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = Path(tmp) / "r.xlsx"
            html_path = Path(tmp) / "r.html"
            self.assertTrue(write_report_xlsx(xlsx, model))
            write_report_html(html_path, model)
            wb = load_workbook(xlsx)
            try:
                headers = [c.value for c in wb["Click Data"][1]]
            finally:
                wb.close()
            html = html_path.read_text(encoding="utf-8")
        self.assertNotIn("Observation", headers)
        self.assertNotIn("Observation flag(s)", headers)
        self.assertEqual(headers[1], "Elapsed time")
        self.assertNotIn("Click Data", html)
        self.assertNotIn("Observation flag(s)", html)
        self.assertIn("state", model["click_data"][0])


class ExcelElapsedTests(unittest.TestCase):
    def test_numeric_duration_under_over_24h(self):
        events = []
        for i, ms in enumerate((500, 3_600_000, 93_784_567)):
            events.append({
                "index": i + 1, "click_index": i + 1, "button": "Left", "accepted": True,
                "event_type": "click", "press_elapsed_ms": ms, "release_elapsed_ms": ms + 8.0,
                "elapsed_ms": ms + 8.0, "press_interval_ms": None if i == 0 else 100.0, "hold_ms": 8.0,
            })
        model = build_completed_result_model(locked=_settings(manual=True), events=events, run_id="dur", target="T")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "d.xlsx"
            self.assertTrue(write_report_xlsx(path, model))
            wb = load_workbook(path)
            try:
                ws = wb["Click Data"]
                for row, ms in ((2, 500), (3, 3_600_000), (4, 93_784_567)):
                    cell = ws.cell(row=row, column=2)
                    self.assertNotIsInstance(cell.value, str)
                    if isinstance(cell.value, timedelta):
                        days = cell.value.total_seconds() / 86_400.0
                    else:
                        self.assertIsInstance(cell.value, (int, float))
                        days = float(cell.value)
                    self.assertTrue(math.isclose(days, ms / 86_400_000.0, rel_tol=0, abs_tol=1e-12))
                    self.assertEqual(cell.number_format, EXCEL_ELAPSED_NUMBER_FORMAT)
                    self.assertIn("[h]", cell.number_format)
            finally:
                wb.close()
        self.assertEqual(format_elapsed_hms(93_784_567), "26:03:04.567")
        self.assertEqual(excel_elapsed_days(3_600_000), 3_600_000 / 86_400_000.0)


class HtmlElapsedTests(unittest.TestCase):
    def test_html_interval_axis_uses_compact_elapsed_not_raw_ms(self):
        events = _events([100.0] * 8)
        events[-1]["press_elapsed_ms"] = 3_600_000.0
        events[-1]["release_elapsed_ms"] = 3_600_001.0
        model = build_completed_result_model(locked=_settings(manual=True), events=events, run_id="html-e", target="T")
        html = render_report_html(model, report_v2.build_report_view(model))
        self.assertNotIn("3600000", html)
        self.assertIn("1:00:00", html)


RUN_DURATION_CASES = (
    (20.1, "00:00:20.100"),
    (3600.0, "01:00:00.000"),
    (93_784.567, "26:03:04.567"),
)


class TestInformationDurationTests(unittest.TestCase):
    def _model(self, seconds):
        return build_completed_result_model(
            locked=_settings(manual=True), events=_events([100.0, 100.0]), run_id="dur-ti",
            target="T", start_time=0.0, stop_time=seconds,
        )

    def test_view_html_and_xlsx_duration_human_readable(self):
        for seconds, expected in RUN_DURATION_CASES:
            with self.subTest(seconds=seconds):
                model = self._model(seconds)
                self.assertEqual(model["run"]["elapsed_s"], seconds)
                view = report_v2.build_report_view(model)
                duration = {key: value for key, _label, value in view["test_information"]}["duration"]
                self.assertEqual(duration, expected)
                html = render_report_html(model, view)
                self.assertIn(expected, html)
                self.assertNotIn(f"{seconds:.1f} s", html)
                with tempfile.TemporaryDirectory() as tmp:
                    path = Path(tmp) / "ti.xlsx"
                    self.assertTrue(write_report_xlsx(path, model))
                    wb = load_workbook(path)
                    try:
                        values = [c.value for row in wb.worksheets[0].iter_rows() for c in row]
                    finally:
                        wb.close()
                self.assertIn(expected, values)
                self.assertNotIn(f"{seconds:.1f} s", values)


class HostInformationCssTests(unittest.TestCase):
    def _css(self):
        def probe():
            return {"ComputerSystem": {"Manufacturer": "Contoso", "Model": "LabHost " + "X" * 120}}

        snap = collect_host_information({}, cim_probe=probe)
        model = build_completed_result_model(
            locked=_settings(manual=True), events=_events([100.0]), run_id="css", target="T",
            host_information=snap,
        )
        html = render_report_html(model, report_v2.build_report_view(model))
        start = html.index("<style>") + len("<style>")
        return html, html[start:html.index("</style>", start)]

    def test_host_values_wrap_after_generic_group_rule(self):
        html, css = self._css()
        generic = ".groups table.kv td{white-space:nowrap;text-align:right}"
        host = ".groups.host-meta table.kv td{white-space:normal;text-align:left;overflow-wrap:anywhere}"
        self.assertIn(generic, css)
        self.assertIn(host, css)
        self.assertGreater(css.index(host), css.index(generic))
        self.assertNotIn("\n.host-meta table.kv td", css)
        self.assertIn("class='groups host-meta'", html)
        self.assertIn("LabHost " + "X" * 120, html)


if __name__ == "__main__":
    unittest.main()
