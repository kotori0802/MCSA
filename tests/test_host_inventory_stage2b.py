"""Stage 2B: frozen host reproducibility snapshot. No live machine BIOS assertions."""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

_SRC = Path(__file__).resolve().parents[1] / "src"
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

from mouse_click_stability_analyzer.app import write_reports
from mouse_click_stability_analyzer.completed_result import build_completed_result_model
from mouse_click_stability_analyzer.host_inventory import (
    HOST_INFORMATION_SCHEMA,
    NOT_AVAILABLE,
    PRIVACY_EXCLUDED_FIELD_NAMES,
    collect_host_information,
    field_display,
    normalize_bios_date,
)
from mouse_click_stability_analyzer.report_v2 import host_information_rows, write_report_html, write_report_xlsx


def _fake_cim():
    return {
        "ComputerSystem": {"Manufacturer": "Contoso", "Model": "LabHost-1"},
        "OperatingSystem": {"Caption": "Windows 11 Pro", "Version": "10.0.26100", "BuildNumber": "26100"},
        "Processor": {"Name": "Fake CPU"},
        "BIOS": {"Manufacturer": "Contoso BIOS", "SMBIOSBIOSVersion": "1.2.3", "ReleaseDate": "20240102000000.000000+000"},
        "USBControllers": [{"Name": "Fake xHCI Host Controller"}],
        "SerialNumber": "SHOULD-NOT-LEAK",
    }


def _events(n: int = 3):
    rows = []
    t = 0.0
    for i in range(n):
        rows.append({
            "index": i + 1, "click_index": i + 1, "button": "Left", "accepted": True,
            "event_type": "click", "press_elapsed_ms": t, "release_elapsed_ms": t + 8.0,
            "elapsed_ms": t + 8.0, "press_interval_ms": None if i == 0 else 100.0, "hold_ms": 8.0,
        })
        t += 100.0
    return rows


class HostInventoryTests(unittest.TestCase):
    def test_injected_cim_populates_core_fields(self):
        snap = collect_host_information({"capture_mode": "Windows Raw Input"}, cim_probe=_fake_cim)
        self.assertEqual(snap["schema"], HOST_INFORMATION_SCHEMA)
        self.assertEqual(snap["host_manufacturer"], "Contoso")
        self.assertEqual(snap["host_model"], "LabHost-1")
        self.assertEqual(snap["os_name"], "Windows 11 Pro")
        self.assertIn("10.0.26100", snap["os_version"])
        self.assertEqual(snap["cpu_model"], "Fake CPU")
        self.assertEqual(snap["bios_vendor"], "Contoso BIOS")
        self.assertEqual(snap["bios_version"], "1.2.3")
        self.assertEqual(snap["bios_date"], "2024-01-02")
        self.assertEqual(snap["usb_host_controllers"], ["Fake xHCI Host Controller"])
        self.assertFalse(snap["dut_usb_controller_attributed"])
        self.assertNotIn("SerialNumber", snap)
        for key in PRIVACY_EXCLUDED_FIELD_NAMES:
            self.assertNotIn(key, snap)

    def test_cim_failure_is_not_available_not_crash(self):
        def boom():
            raise RuntimeError("cim failed")

        snap = collect_host_information({}, cim_probe=boom)
        self.assertEqual(field_display(snap, "host_manufacturer"), NOT_AVAILABLE)
        self.assertEqual(field_display(snap, "bios_version"), NOT_AVAILABLE)
        self.assertTrue(snap["os_name"])
        self.assertTrue(snap["capture_backend"])
        self.assertTrue(snap["tool_version"])

    def test_snapshot_frozen_once_html_xlsx_match(self):
        snap = collect_host_information({"capture_mode": "Windows Raw Input"}, cim_probe=_fake_cim)
        settings = {
            "manual_mode": True, "focus_button": "Left", "profile_name": "Manual",
            "capture_mode": "Windows Raw Input", "fixture_reference_hz": 0,
        }
        model = build_completed_result_model(
            locked=settings, events=_events(), run_id="host1", target="DUT",
            host_information=snap,
        )
        self.assertIsNotNone(model["run"]["host_information"])
        self.assertEqual(model["run"]["host_information"]["bios_version"], "1.2.3")
        mutated = dict(snap)
        mutated["bios_version"] = "CHANGED"
        self.assertEqual(model["run"]["host_information"]["bios_version"], "1.2.3")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            paths = write_reports(root, "DUT", "", settings, _events(), "host1", completed_model=model)
            html_text = Path(paths["engineering_html"]).read_text(encoding="utf-8")
            from openpyxl import load_workbook
            wb = load_workbook(paths["engineering_xlsx"])
            try:
                xlsx_text = "\n".join(
                    str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None
                )
            finally:
                wb.close()
        self.assertIn("Contoso BIOS", html_text)
        self.assertIn("Contoso BIOS", xlsx_text)
        self.assertIn("Fake xHCI Host Controller", html_text)
        self.assertIn("1.2.3", html_text)
        self.assertNotIn("SHOULD-NOT-LEAK", html_text)
        self.assertNotIn("SHOULD-NOT-LEAK", xlsx_text)

    def test_write_reports_does_not_requery_when_frozen(self):
        snap = collect_host_information({}, cim_probe=_fake_cim)
        settings = {"manual_mode": True, "focus_button": "Left", "profile_name": "Manual"}
        model = build_completed_result_model(
            locked=settings, events=_events(), run_id="host2", target="DUT", host_information=snap,
        )
        probe = mock.Mock(return_value=_fake_cim())
        with tempfile.TemporaryDirectory() as tmp, mock.patch(
            "mouse_click_stability_analyzer.app.collect_host_information",
            side_effect=AssertionError("must not re-collect"),
        ):
            write_reports(Path(tmp), "DUT", "", settings, _events(), "host2", completed_model=model)
        probe.assert_not_called()

    def test_privacy_fields_never_emitted(self):
        rows = host_information_rows(
            collect_host_information({}, cim_probe=_fake_cim), tool_version="1.0.0",
        )
        blob = " ".join(f"{i} {l} {v}" for i, l, v in rows).lower()
        for token in ("serial", "username", "mac address", "ip address", "product key", "password"):
            self.assertNotIn(token, blob)

    def test_collector_failure_does_not_block_report(self):
        settings = {"manual_mode": True, "focus_button": "Left", "profile_name": "Manual"}
        model = build_completed_result_model(locked=settings, events=_events(), run_id="host3", target="DUT")
        with tempfile.TemporaryDirectory() as tmp, mock.patch(
            "mouse_click_stability_analyzer.app.collect_host_information",
            side_effect=RuntimeError("inventory failed"),
        ):
            paths = write_reports(Path(tmp), "DUT", "", settings, _events(), "host3", completed_model=model)
            self.assertTrue(Path(paths["engineering_html"]).exists())
            html = Path(paths["engineering_html"]).read_text(encoding="utf-8")
            self.assertIn("Not available", html)


class BiosDateNormalizationTests(unittest.TestCase):
    def test_supported_collector_forms(self):
        self.assertEqual(normalize_bios_date("20240102000000.000000+000"), "2024-01-02")
        self.assertEqual(normalize_bios_date("/Date(1773792000000)/"), "2026-03-18")
        self.assertEqual(normalize_bios_date(r"\/Date(1773792000000)\/"), "2026-03-18")
        self.assertEqual(normalize_bios_date("/Date(1773792000000+0800)/"), "2026-03-18")
        self.assertEqual(normalize_bios_date("2024-01-02"), "2024-01-02")

    def test_leap_year_calendar(self):
        self.assertEqual(normalize_bios_date("20240229000000.000000+000"), "2024-02-29")
        self.assertIsNone(normalize_bios_date("20230229000000.000000+000"))

    def test_malformed_and_impossible_are_unavailable(self):
        self.assertIsNone(normalize_bios_date(None))
        self.assertIsNone(normalize_bios_date(""))
        self.assertIsNone(normalize_bios_date("   "))
        self.assertIsNone(normalize_bios_date("not-a-date"))
        self.assertIsNone(normalize_bios_date("1773-79-20"))
        self.assertIsNone(normalize_bios_date("2024-79-20"))
        self.assertIsNone(normalize_bios_date("2024-01-99"))
        self.assertIsNone(normalize_bios_date("202603**********.******+***"))
        self.assertIsNone(normalize_bios_date("********000000.000000+000"))
        self.assertIsNone(normalize_bios_date("2026-03-18T08:00:00+08:00"))
        self.assertIsNone(normalize_bios_date("/Date(not-ms)/"))
        self.assertEqual(field_display({"bios_date": normalize_bios_date("1773-79-20")}, "bios_date"), NOT_AVAILABLE)

    def test_invalid_raw_never_reaches_html_xlsx(self):
        def probe():
            payload = _fake_cim()
            payload["BIOS"] = dict(payload["BIOS"], ReleaseDate="/Date(1773792000000)/X")
            return payload

        snap = collect_host_information({}, cim_probe=probe)
        self.assertIsNone(snap["bios_date"])
        settings = {"manual_mode": True, "focus_button": "Left", "profile_name": "Manual"}
        model = build_completed_result_model(
            locked=settings, events=_events(), run_id="bios-bad", target="DUT", host_information=snap,
        )
        with tempfile.TemporaryDirectory() as tmp:
            html_path = Path(tmp) / "b.html"
            xlsx_path = Path(tmp) / "b.xlsx"
            write_report_html(html_path, model)
            write_report_xlsx(xlsx_path, model)
            html = html_path.read_text(encoding="utf-8")
            from openpyxl import load_workbook
            wb = load_workbook(xlsx_path)
            try:
                xlsx = "\n".join(
                    str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None
                )
            finally:
                wb.close()
        self.assertNotIn("1773-79-20", html)
        self.assertNotIn("1773-79-20", xlsx)
        self.assertIn("Not available", html)

    def test_valid_normalized_date_identical_in_html_xlsx(self):
        def probe():
            payload = _fake_cim()
            payload["BIOS"] = dict(payload["BIOS"], ReleaseDate="/Date(1773792000000)/")
            return payload

        snap = collect_host_information({}, cim_probe=probe)
        self.assertEqual(snap["bios_date"], "2026-03-18")
        settings = {"manual_mode": True, "focus_button": "Left", "profile_name": "Manual"}
        model = build_completed_result_model(
            locked=settings, events=_events(), run_id="bios-ok", target="DUT", host_information=snap,
        )
        with tempfile.TemporaryDirectory() as tmp:
            html_path = Path(tmp) / "b.html"
            xlsx_path = Path(tmp) / "b.xlsx"
            write_report_html(html_path, model)
            write_report_xlsx(xlsx_path, model)
            html = html_path.read_text(encoding="utf-8")
            from openpyxl import load_workbook
            wb = load_workbook(xlsx_path)
            try:
                xlsx = "\n".join(
                    str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None
                )
            finally:
                wb.close()
        self.assertIn("2026-03-18", html)
        self.assertIn("2026-03-18", xlsx)
        self.assertEqual(html.count("2026-03-18"), 1)


if __name__ == "__main__":
    unittest.main()
