"""Phase 4D / 4D-1 / 4D-2: engineering HTML/XLSX are factual records rendered from the canonical model."""
from __future__ import annotations

import copy
import html as html_lib
import json
import os
import re
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from unittest import mock

from PySide6.QtCore import QCoreApplication, QEvent

_ROOT = Path(__file__).resolve().parents[1]
_SRC = _ROOT / "src"
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))
if str(_ROOT / "tests") not in sys.path:
    sys.path.insert(0, str(_ROOT / "tests"))

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook

from mouse_click_stability_analyzer import app as legacy_app
from mouse_click_stability_analyzer import report_v2
from mouse_click_stability_analyzer.completed_result import build_completed_result_model
from mouse_click_stability_analyzer.overlay_io import extract_overlay_events, load_overlay_payload
from association_support import apply_explicit_association
from mouse_click_stability_analyzer.qt_shell import app as qt_app

_RETIRED = ("MISS_CLICK_WARN", "MISS_CLICK_FAIL", "EXCESS_CLICK_COUNT_WARN", "EXCESS_CLICK_COUNT_FAIL")
_EXPLANATORY = (
    "Observations",
    "Anomaly Focus",
    "No current rule-based observations",
    "No individual Host Press",
    "Per-event timing error is not charted",
    "Host-observed schedule association only",
    "Host Evidence is not applicable to Manual mode",
    "not applicable",
    "Companion files",
    "Existing rule",
    "summary.json",
    "detail.json",
    "events.csv",
    "analysis.txt",
    "report.xlsx",
)
_LEGACY = (
    "Expected / Actual clicks",
    "Missed / Excess click count",
    "Actual clicks",
    "Root Cause Hints",
    "Expected time",
    "Drift",
    "Within ±0.5",
    "Showing first 80",
    "Test Method",
    "Detail Paths",
    "Run Details",
    "Assessment",
    "Issues",
)
_HOST_EVIDENCE_LABELS = (
    "Schedule windows",
    "Configured reference interval",
    "Schedule Association Window",
    "Host observations in capture",
    "Host observations recorded",
    "Host observations − schedule windows",
    "Schedule windows without matched Host observation",
    "Schedule windows with 1 matched Host observation",
    "Schedule windows with >1 matched Host observations",
    "Host observations matched to schedule",
    "Host observations not matched to schedule",
    "Ambiguous Host observations",
    "Host observations outside capture",
    "Matched observations used",
    "Average matched timing error",
    "Average absolute matched timing error",
    "Minimum matched timing error",
    "Maximum matched timing error",
    "Matched timing error stdev",
)
_RETIRED_HOST_EVIDENCE_LABELS = (
    "Unoccupied windows", "Single-occupancy windows", "Multi-occupancy windows", "Matched", "Unmatched",
    "Ambiguous", "Outside capture", "Observation count delta", "Raw Host observations", "Stdev",
    "Average matched error", "Tolerance", "Reference interval", "Configured tolerance",
    "Host observations recorded (all)", "Host observations - schedule windows",
)
_PHYSICAL_CLAIM_CODES = (
    "MISS_OR_SLOW", "HIGH_PROBABILITY_DOUBLE_CLICK", "TRUE_EXTRA_CLICK", "INCOMPLETE_REBOUND",
    "FIXTURE_ACTUATION_TOO_SHORT", "SETUP_ALIGNMENT_SUSPECT", "FIXTURE_HZ_MISMATCH",
)
_CHART_TITLES = (
    "Elapsed time vs Click interval",
    "Click interval distribution",
    "Click interval vs Hold duration",
)
_REF_MS = 1000.0 / 5.6
_GOLDEN_STEP_MS = 174.0
_MANUAL_OFFSETS = [0, 300, 610, 905, 1220, 1500, 1790, 2120, 2400, 2715]


def _golden_offsets() -> list[float]:
    """58 Host Presses at ~5.75 Hz against a 5.6 Hz reference, with one early/late timing pair.

    Legacy expected-count promotion sees 56 windows and 58 presses (EXCESS_CLICK_COUNT_FAIL);
    the canonical model sees only TIMING_JITTER_PAIR.
    """
    intervals = [_GOLDEN_STEP_MS] * 57
    intervals[20] = 150.0
    intervals[21] = 2 * _REF_MS - 150.0
    offsets = [0.0]
    for step in intervals:
        offsets.append(offsets[-1] + step)
    return offsets


def _visible_html_text(text: str) -> str:
    text = re.sub(r"<style>.*?</style>", " ", text, flags=re.S)
    text = re.sub(r"<title>.*?</title>", " ", text, flags=re.S)
    text = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", text)


def _html_kv(text: str, section_id: str) -> dict[str, str]:
    match = re.search(rf"<section id=['\"]{section_id}['\"]>(.*?)</section>", text, flags=re.S)
    if not match:
        return {}
    pairs = re.findall(r"<th scope='row'>(.*?)</th><td[^>]*>(.*?)</td>", match.group(1), flags=re.S)
    return {
        html_lib.unescape(re.sub(r"<[^>]+>", "", k)): html_lib.unescape(re.sub(r"<[^>]+>", "", v))
        for k, v in pairs
    }


def _html_run_status(text: str) -> dict[str, str]:
    return _html_kv(text, "run-status")


def _sheet_kv(path: Path, sheet: str) -> dict[str, object]:
    wb = load_workbook(path)
    try:
        out: dict[str, object] = {}
        for row in wb[sheet].iter_rows(min_col=1, max_col=2, values_only=True):
            if row[0] is not None and row[0] not in out:
                out[row[0]] = row[1]
        return out
    finally:
        wb.close()


def _xlsx_run_status(path: Path) -> str:
    wb = load_workbook(path)
    try:
        values = [row[1] for row in wb["Summary"].iter_rows(min_col=1, max_col=2, values_only=True)
                  if row[0] == "Run Status" and row[1] is not None]
    finally:
        wb.close()
    assert len(values) == 1, values
    return values[0]


_VERDICT_WORDS = ("PASS", "WARN", "FAIL")
_RETIRED_RESULT_LABELS = ("Test Result", "Result Code", "Result Codes")


def _assert_no_verdict_words(test: unittest.TestCase, text: str) -> None:
    for word in _VERDICT_WORDS:
        test.assertIsNone(re.search(rf"\b{word}\b", text), word)
    for label in _RETIRED_RESULT_LABELS:
        test.assertNotIn(label, text, label)


def _workbook_text(path: Path) -> str:
    wb = load_workbook(path)
    try:
        parts = list(wb.sheetnames)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                parts.extend(str(v) for v in row if v is not None)
        return "\n".join(parts)
    finally:
        wb.close()


def _sheet_names(path: Path) -> tuple[str, ...]:
    wb = load_workbook(path)
    try:
        return tuple(wb.sheetnames)
    finally:
        wb.close()


_C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _xlsx_charts(path: Path) -> dict[str, list[dict[str, str]]]:
    """Chart title -> series [{title, x, y}] read from the saved chart XML."""
    out: dict[str, list[dict[str, str]]] = {}
    with zipfile.ZipFile(path) as zf:
        for name in sorted(n for n in zf.namelist() if n.startswith("xl/charts/chart")):
            root = ET.fromstring(zf.read(name))
            chart = root.find(f"{_C}chart")
            title = "".join(t.text or "" for t in chart.find(f"{_C}title").iter(f"{_A}t"))
            series = []
            for ser in chart.iter(f"{_C}ser"):
                def ref(*tags):
                    for tag in tags:
                        node = ser.find(f"{_C}{tag}")
                        if node is not None:
                            f = node.find(f".//{_C}f")
                            return f.text if f is not None else ""
                    return ""
                tx = ser.find(f"{_C}tx")
                label = ""
                if tx is not None:
                    v = tx.find(f".//{_C}v")
                    label = v.text if v is not None else (tx.find(f".//{_C}f").text or "")
                series.append({"title": label, "x": ref("xVal", "cat"), "y": ref("yVal", "val")})
            out[title] = series
    return out


def _html_figures(text: str) -> dict[str, str]:
    return {
        caption: body
        for caption, body in re.findall(r"<figure><figcaption>(.*?)</figcaption>(.*?)</figure>", text, flags=re.S)
    }


def _make_qapp():
    from PySide6.QtWidgets import QApplication

    return QApplication.instance() or QApplication([])


class ReportV2QtTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._qapp = _make_qapp()
        cls._tmp = tempfile.TemporaryDirectory()

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def setUp(self):
        self.root = Path(self._tmp.name) / self.id().rsplit(".", 1)[-1]
        self.root.mkdir()
        self.win = qt_app.MainWindow(report_root=self.root)
        self.win._listen_start_timer.stop()
        self.win.ui_timer.stop()
        self.win.ensure_raw_input = lambda: None
        self.win._cursor_in_test_area = lambda: True
        self.win._cursor_over_capture_exclusion_ui = lambda: False
        self.win.listening = True

    def tearDown(self):
        self.win.close()
        self.win.deleteLater()
        self._qapp.processEvents()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)

    def _run(self, offsets, *, manual=False, hz="5.6", auto_sec="10", device="123"):
        self.win.device_edit.setText(device)
        if manual:
            self.win.radio_manual.setChecked(True)
        else:
            self.win.radio_fixture.setChecked(True)
            self.win.hz_combo.setCurrentText(hz)
            self.win.timing_tol_edit.setText("2")
            if auto_sec is None:
                self.win.auto_stop_off.setChecked(True)
            else:
                self.win.auto_stop_on.setChecked(True)
                self.win.auto_stop_sec_combo.setCurrentText(auto_sec)
        self.assertTrue(self.win.lock_profile())
        if not manual:
            apply_explicit_association(self.win, 2.0)
        self.assertTrue(self.win.arm_capture())
        t0 = float(self.win._arm_boundary_perf) + 0.001
        self.win._begin_session_from_first_click("Left", t0, t0 + 0.04, {})
        for off in offsets[1:]:
            press = t0 + off / 1000.0
            self.win.record_click("Left", press, press + 0.04, {})
        self.win.stop_run()
        if self.win._post_run_input_guard_active:
            self.win._release_post_run_input_guard()
        self.assertTrue(self.win.completed)

    def _generate(self, preferred="engineering_html") -> dict[str, str]:
        with mock.patch.object(qt_app.QMessageBox, "information"), \
             mock.patch.object(qt_app.QMessageBox, "critical") as critical:
            self.win.finish_report(preferred_artifact=preferred)
        critical.assert_not_called()
        self.assertTrue(self.win.report_written)
        return dict(self.win.report_paths)

    def _artifacts(self, paths):
        return Path(paths["engineering_html"]).read_text(encoding="utf-8"), Path(paths["engineering_xlsx"])

    def _golden(self):
        self._run(_golden_offsets())
        facts = self.win._completed_run_facts()
        return facts, self._generate()

    def _manual(self):
        self._run(_MANUAL_OFFSETS, manual=True)
        facts = self.win._completed_run_facts()
        return facts, self._generate()

    # --- golden regression -------------------------------------------------

    def test_golden_fixture_run_is_complete_everywhere(self):
        facts, paths = self._golden()
        self.assertEqual(facts["run_status"], "COMPLETE")
        self.assertNotIn("verdict", facts)
        self.assertEqual(facts["tags"], ["TIMING_JITTER_PAIR"])
        evidence = json.loads(Path(paths["evidence_json"]).read_text(encoding="utf-8"))
        self.assertEqual(evidence["run_status"], "COMPLETE")
        self.assertNotIn("result", evidence)
        self.assertNotIn("legacy_heuristic_assessment", evidence)
        self.assertNotIn("issue_tags", evidence)
        html_text, xlsx = self._artifacts(paths)
        self.assertEqual(_html_run_status(html_text), {"Run Status": "COMPLETE"})
        self.assertEqual(_xlsx_run_status(xlsx), "COMPLETE")
        visible = _visible_html_text(html_text)
        workbook = _workbook_text(xlsx)
        for text in (visible, workbook):
            _assert_no_verdict_words(self, text)
            self.assertNotIn("EXCESS", text)
            self.assertNotIn("MISS_CLICK", text)
        info = _html_kv(html_text, "test-information")
        self.assertEqual(info["Device"], "123")
        self.assertEqual(info["Run Profile"], "Fixture_Custom")
        self.assertEqual(info["Analysis scope"], "Measurement only")
        self.assertNotIn("Screening", info)
        self.assertEqual(info["Reference Rate"], "5.6 Hz")
        self.assertEqual(info["Mode"], "Fixture")

    def test_html_and_xlsx_run_status_equal_canonical(self):
        facts, paths = self._golden()
        html_text, xlsx = self._artifacts(paths)
        expected = facts["model"]["run"]["run_status"]
        self.assertEqual(expected, "COMPLETE")
        self.assertEqual(_html_run_status(html_text)["Run Status"], expected)
        self.assertEqual(_xlsx_run_status(xlsx), expected)
        self.assertIn("<section id=\"run-status\"><h2>Run Status</h2>", html_text)
        self.assertNotIn("test-result", html_text)

    def test_report_host_presses_matches_canonical_completed_count(self):
        facts, paths = self._golden()
        canonical = int(facts["model"]["run"]["host_press_count"])
        self.assertEqual(canonical, len(_golden_offsets()))
        self.assertEqual(int(self.win.val_clicks.text()), canonical)
        html_text, xlsx = self._artifacts(paths)
        info = _html_kv(html_text, "test-information")
        self.assertEqual(info["Host Presses"], str(canonical))
        summary = _sheet_kv(xlsx, "Summary")
        self.assertEqual(int(summary["Host Presses"]), canonical)
        self.assertNotIn("waiting_for_mouse_input", html_text)
        self.assertNotIn("Waiting for mouse input", _visible_html_text(html_text))

    def test_retired_tags_absent_from_html_and_xlsx(self):
        _, paths = self._golden()
        html_text, xlsx = self._artifacts(paths)
        workbook = _workbook_text(xlsx)
        for tag in _RETIRED:
            self.assertNotIn(tag, html_text)
            self.assertNotIn(tag, workbook)

    def test_measurements_match_canonical_metrics(self):
        facts, paths = self._golden()
        metrics = facts["model"]["assessment"]["metrics"]
        html_text, xlsx = self._artifacts(paths)
        avg = metrics["interval_avg_ms"]
        delta = avg - _REF_MS
        expected = {
            "Host-observed average interval": f"{avg:.2f} ms",
            "Host-observed cadence": f"{1000.0 / avg:.2f} Hz",
            "Reference Interval": f"{_REF_MS:.2f} ms",
            "Average interval Δ vs reference": f"{delta:+.2f} ms" if delta else f"{delta:.2f} ms",
            "Interval CV": f"{metrics['interval_cv_pct']:.2f} %",
            "Average hold": f"{metrics['hold_avg_ms']:.2f} ms",
        }
        html_kv = _html_kv(html_text, "measurements")
        xlsx_kv = _sheet_kv(xlsx, "Measurements")
        for label, value in expected.items():
            self.assertEqual(html_kv[label], value, label)
            self.assertEqual(xlsx_kv[label], value, label)
        self.assertEqual(html_kv["Host Press count"], "58")
        self.assertEqual(xlsx_kv["Host Press count"], 58)
        self.assertIn("Elapsed time vs Click interval", html_text)
        self.assertIn("Click interval distribution", html_text)
        self.assertIn("Click interval vs Hold duration", html_text)

    def test_host_evidence_numbers_match_snapshot_without_prose(self):
        facts, paths = self._golden()
        snap = facts["model"]["host_evidence"]["snapshot"]
        html_text, xlsx = self._artifacts(paths)
        counts, matching = snap["counts"], snap["matching"]
        expected = {
            "Schedule windows": str(counts["expected_slot_count"]),
            "Host observations in capture": str(counts["in_window_host_press_count"]),
            "Host observations recorded": str(counts["raw_host_press_count"]),
            "Host observations − schedule windows": f"{counts['host_count_delta_vs_expected_slots']:+d}",
            "Schedule windows without matched Host observation": str(matching["empty_expected_slot_count"]),
            "Schedule windows with 1 matched Host observation": str(matching["single_occupancy_slot_count"]),
            "Schedule windows with >1 matched Host observations": str(matching["multi_occupancy_slot_count"]),
            "Host observations matched to schedule": str(matching["matched_host_observation_count"]),
            "Host observations not matched to schedule": str(matching["unmatched_host_observation_count"]),
            "Ambiguous Host observations": str(matching["ambiguous_host_observation_count"]),
            "Host observations outside capture": str(matching["outside_capture_window_observation_count"]),
            "Matched observations used": f"{matching['matched_host_observation_count']} of {counts['in_window_host_press_count']} in capture",
        }
        self.assertEqual(expected["Host observations − schedule windows"], "+2")
        self.assertEqual(
            int(expected["Host observations − schedule windows"]),
            counts["in_window_host_press_count"] - counts["expected_slot_count"],
        )
        html_kv = _html_kv(html_text, "host-evidence")
        xlsx_kv = _sheet_kv(xlsx, "Host Evidence")
        for label, value in expected.items():
            self.assertEqual(html_kv[label], value, label)
            self.assertEqual(str(xlsx_kv[label]), value, label)
        for label in ("Average matched timing error", "Average absolute matched timing error",
                      "Minimum matched timing error", "Maximum matched timing error", "Matched timing error stdev"):
            self.assertTrue(html_kv[label].endswith(" ms"), label)

    def test_host_evidence_labels_are_explicit_schedule_and_host_wording(self):
        _, paths = self._golden()
        html_text, xlsx = self._artifacts(paths)
        html_labels = list(_html_kv(html_text, "host-evidence"))
        self.assertEqual(html_labels, list(_HOST_EVIDENCE_LABELS))
        xlsx_kv = _sheet_kv(xlsx, "Host Evidence")
        for label in _HOST_EVIDENCE_LABELS:
            self.assertIn(label, xlsx_kv, label)
        for retired in _RETIRED_HOST_EVIDENCE_LABELS:
            self.assertNotIn(retired, html_labels, retired)
            self.assertNotIn(retired, xlsx_kv, retired)
        section = re.search(r"<section id='host-evidence'>(.*?)</section>", html_text, flags=re.S).group(1)
        self.assertNotIn("<p", section)

    def test_fixture_html_charts_identify_host_series_and_configured_reference(self):
        _, paths = self._golden()
        html_text, _ = self._artifacts(paths)
        figures = _html_figures(html_text)
        captions = tuple(figures)
        self.assertEqual(captions[0], "Elapsed time vs Click interval")
        self.assertTrue(captions[1].startswith("Click interval distribution (N="), captions[1])
        self.assertEqual(captions[2], "Elapsed time vs Hold duration")
        self.assertEqual(captions[3], "Click interval vs Hold duration")
        hold_body = figures.pop("Elapsed time vs Hold duration")
        for neutral in ("class='ref'", "class='band'", "class='bnd'", "class='out'", "class='legend'"):
            self.assertNotIn(neutral, hold_body)
        self.assertIn("Hold duration (ms)", hold_body)
        for title, body in figures.items():
            with self.subTest(chart=title):
                self.assertEqual(body.count("class='ref'"), 1)
                self.assertNotIn("class='tol'", body)
                self.assertNotIn("Association", body)
                self.assertIn(f"class='rl' text-anchor='", body)
                self.assertIn(f">{_REF_MS:.2f} ms</text>", body)
                legend = re.findall(r"<span class='li'>.*?</span>", body)
                self.assertEqual(
                    [re.sub(r"<[^>]+>", "", item) for item in legend],
                    ["Host observation", f"Reference interval {_REF_MS:.2f} ms"],
                )

    def test_run_context_uses_schedule_association_window_wording(self):
        _, paths = self._golden()
        html_text, xlsx = self._artifacts(paths)
        info = _html_kv(html_text, "test-information")
        summary = _sheet_kv(xlsx, "Summary")
        relative_label = "Window relative to reference interval (derived, informational)"
        for kv in (info, summary):
            self.assertEqual(kv["Schedule Association Window"], "±2.00 ms")
            self.assertEqual(kv[relative_label], f"{2.0 / _REF_MS * 100.0:.1f} %")
        visible = _visible_html_text(html_text)
        workbook = _workbook_text(xlsx)
        for text in (visible, workbook):
            _assert_no_verdict_words(self, text)
            for stale in ("Fixture Timing Tolerance", "Timing Tolerance", "Configured association tolerance",
                          "Association window"):
                self.assertNotIn(stale, text)

    def test_manual_html_charts_have_no_configured_reference_marker(self):
        _, paths = self._manual()
        html_text, xlsx = self._artifacts(paths)
        figures = _html_figures(html_text)
        captions = tuple(figures)
        self.assertEqual(captions[0], "Elapsed time vs Click interval")
        self.assertTrue(captions[1].startswith("Click interval distribution (N="), captions[1])
        self.assertEqual(captions[2], "Elapsed time vs Hold duration")
        self.assertEqual(captions[3], "Click interval vs Hold duration")
        for title, body in figures.items():
            with self.subTest(chart=title):
                self.assertNotIn("class='ref'", body)
                self.assertNotIn("class='tol'", body)
                self.assertNotIn("class='rl'", body)
                self.assertNotIn("class='legend'", body)
        self.assertNotIn("Configured reference", html_text)
        self.assertNotIn("Configured tolerance", html_text)
        self.assertNotIn("Association Window", html_text)
        for title, series in _xlsx_charts(xlsx).items():
            with self.subTest(xlsx_chart=title):
                self.assertEqual(len(series), 1)
                self.assertNotIn("Configured", series[0]["title"])

    def test_xlsx_charts_keep_axis_sources_and_add_reference_series(self):
        _, paths = self._golden()
        charts = _xlsx_charts(Path(paths["engineering_xlsx"]))
        hold_over_time = charts.pop("Elapsed time vs Hold duration")
        self.assertEqual([s["title"] for s in hold_over_time], ["Host observation"])
        self.assertRegex(hold_over_time[0]["x"], r"^'Click Data'!\$B\$2:\$B\$\d+$")
        self.assertRegex(hold_over_time[0]["y"], r"^'Click Data'!\$E\$2:\$E\$\d+$")
        self.assertEqual(set(charts), set(_CHART_TITLES))
        interval, hold = charts["Elapsed time vs Click interval"][0], charts["Click interval vs Hold duration"][0]
        self.assertRegex(interval["x"], r"^'Click Data'!\$B\$2:\$B\$\d+$")
        self.assertRegex(interval["y"], r"^'Click Data'!\$D\$2:\$D\$\d+$")
        self.assertRegex(hold["x"], r"^'Click Data'!\$D\$2:\$D\$\d+$")
        self.assertRegex(hold["y"], r"^'Click Data'!\$E\$2:\$E\$\d+$")
        histogram = charts["Click interval distribution"][0]
        self.assertTrue(histogram["x"].startswith("'Measurements'!$A$"))
        self.assertTrue(histogram["y"].startswith("'Measurements'!$D$"))
        for title in _CHART_TITLES:
            with self.subTest(chart=title):
                self.assertEqual(
                    [s["title"] for s in charts[title]],
                    ["Host observation", f"Reference interval {_REF_MS:.2f} ms"],
                )
        wb = load_workbook(paths["engineering_xlsx"])
        try:
            ws = wb["Measurements"]
            self.assertTrue(all(ws.column_dimensions[c].hidden for c in ("AA", "AB", "AF", "AJ")))
        finally:
            wb.close()

    def test_unavailable_dut_identity_fields_are_omitted(self):
        _, paths = self._golden()
        html_text, xlsx = self._artifacts(paths)
        info = _html_kv(html_text, "test-information")
        summary = _sheet_kv(xlsx, "Summary")
        for label in ("HID device name", "VID", "PID", "bcdDevice", "Firmware", "Connection", "Operator",
                      "DUT", "DUT Sample ID", "Lab Host"):
            self.assertNotIn(label, info, label)
        for label in ("HID device name", "VID", "PID", "bcdDevice", "Connection", "Operator",
                      "DUT", "DUT Sample ID", "Lab Host"):
            self.assertNotIn(label, summary, label)
        self.assertNotIn("—", "".join(info.values()))
        self.assertNotIn(None, [summary[k] for k in info])

    def test_fixture_reports_contain_no_explanatory_or_artifact_content(self):
        _, paths = self._golden()
        html_text, xlsx = self._artifacts(paths)
        visible = _visible_html_text(html_text)
        workbook = _workbook_text(xlsx)
        for forbidden in _EXPLANATORY + _LEGACY:
            self.assertNotIn(forbidden, visible, forbidden)
            self.assertNotIn(forbidden, workbook, forbidden)
        self.assertNotIn("<a ", html_text)
        self.assertNotIn("href", html_text)
        for key in ("evidence_json", "all_click_log_csv", "engineering_xlsx"):
            self.assertNotIn(Path(paths[key]).name, html_text)
            self.assertNotIn(Path(paths[key]).name, workbook)

    def test_manual_report_is_complete_without_result_code_or_host_evidence(self):
        facts, paths = self._manual()
        self.assertEqual(facts["run_status"], "COMPLETE")
        self.assertIn("MANUAL_PROFILE_NO_EXPECTED_COUNT", facts["tags"])
        html_text, xlsx = self._artifacts(paths)
        self.assertEqual(_html_run_status(html_text), {"Run Status": "COMPLETE"})
        self.assertEqual(_xlsx_run_status(xlsx), "COMPLETE")
        visible = _visible_html_text(html_text)
        workbook = _workbook_text(xlsx)
        for text in (visible, workbook):
            _assert_no_verdict_words(self, text)
            self.assertNotIn("MANUAL_PROFILE_NO_EXPECTED_COUNT", text)
            self.assertNotIn("Host Evidence", text)
            for absent in ("Schedule windows", "Reference Rate", "Schedule Association Window",
                           "Fixture Timing Tolerance", "Interval Δ vs ref."):
                self.assertNotIn(absent, text)
            for forbidden in _EXPLANATORY + _LEGACY:
                self.assertNotIn(forbidden, text, forbidden)
        self.assertNotIn("host-evidence", html_text)
        self.assertEqual(_sheet_names(xlsx), ("Summary", "Measurements", "Click Data"))
        self.assertIn("Measurements", visible)
        self.assertEqual(_html_kv(html_text, "test-information")["Mode"], "Manual")

    def test_fixture_without_snapshot_omits_host_evidence(self):
        self._run(_golden_offsets()[:20], auto_sec=None)
        self.assertIsNone(self.win.host_evidence_snapshot)
        html_text, xlsx = self._artifacts(self._generate())
        self.assertNotIn("Host Evidence", _visible_html_text(html_text))
        self.assertEqual(_sheet_names(xlsx), ("Summary", "Measurements", "Click Data"))

    def test_fixture_xlsx_sheet_set_is_exact(self):
        _, paths = self._golden()
        self.assertEqual(_sheet_names(Path(paths["engineering_xlsx"])),
                         ("Summary", "Measurements", "Host Evidence", "Click Data"))

    def test_click_data_xlsx_full_table_and_html_omits_it(self):
        facts, paths = self._golden()
        html_text, xlsx = self._artifacts(paths)
        wb = load_workbook(xlsx)
        try:
            ws = wb["Click Data"]
            headers = [c.value for c in ws[1]]
            self.assertEqual(headers, report_v2.click_data_columns(False))
            self.assertNotIn("Observation", headers)
            self.assertNotIn("Observation flag(s)", headers)
            self.assertNotIn("Rule Code(s)", headers)
            self.assertNotIn("State", headers)
            self.assertEqual(ws.freeze_panes, "A2")
            self.assertTrue(ws.auto_filter.ref)
            self.assertEqual(ws.max_row - 1, len(facts["model"]["click_data"]))
        finally:
            wb.close()
        for forbidden in ("Expected", "Drift", "raw_hdevice", "raw_dx", "raw_dy", "flags", "Tags"):
            self.assertFalse(any(forbidden.lower() == str(h).lower() or forbidden in str(h) for h in headers), forbidden)
        visible = _visible_html_text(html_text)
        self.assertNotIn("Click Data", visible)
        self.assertNotIn("Rule Code(s)", visible)
        self.assertNotIn("ON_CADENCE", visible)

    def test_host_information_facts_only(self):
        _, paths = self._golden()
        html_text, xlsx = self._artifacts(paths)
        html_kv = _html_kv(html_text, "host-information")
        xlsx_kv = _sheet_kv(xlsx, "Summary")
        for label in ("Capture backend", "Tool version"):
            self.assertTrue(html_kv.get(label), label)
            self.assertTrue(xlsx_kv.get(label), label)
        self.assertEqual(html_kv["Tool version"], legacy_app.TOOL_VERSION)
        self.assertIn("USB host controllers", html_kv)
        self.assertIn("Python runtime", html_kv)

    def test_reports_have_no_absolute_paths(self):
        _, paths = self._golden()
        html_text, xlsx = self._artifacts(paths)
        workbook = _workbook_text(xlsx)
        for text in (html_text, workbook):
            self.assertNotIn("file://", text)
            self.assertNotIn(str(self.root), text)
            self.assertIsNone(re.search(r"[A-Za-z]:\\", text))

    # --- artifact contract --------------------------------------------------

    def test_generate_buttons_still_write_full_bundle(self):
        self._run(_golden_offsets())
        self.win.report_written = False
        with mock.patch.object(qt_app.QMessageBox, "information") as info, \
             mock.patch.object(qt_app.QMessageBox, "critical") as critical:
            self.win.finish_report()
        critical.assert_not_called()
        info.assert_called_once()
        self.assertEqual(
            info.call_args[0][2],
            self.win._tr("saved_to", path=Path(self.win.report_root) / f"mcsa_{self.win.run_id}"),
        )
        self.assertTrue(legacy_app._verify_artifacts(self.win.report_paths)["ok"])
        for key, fname in legacy_app.artifact_names(self.win.run_id).items():
            self.assertTrue(Path(self.win.report_paths[key]).is_file(), fname)

    def test_direct_secondary_export_is_silent_with_real_bundle(self):
        self._run(_golden_offsets())
        dest = self.root / "exported.html"
        with mock.patch.object(qt_app.QMessageBox, "information") as info, \
             mock.patch.object(qt_app.QMessageBox, "critical") as critical, \
             mock.patch.object(qt_app.QFileDialog, "getSaveFileName", return_value=(str(dest), "")):
            self.win.export_report_artifact("engineering_html")
        info.assert_not_called()
        critical.assert_not_called()
        self.assertTrue(self.win.report_written)
        self.assertIn(report_v2.REPORT_TITLE, dest.read_text(encoding="utf-8"))

    def test_save_report_as_copies_full_bundle(self):
        self._run(_golden_offsets())
        out_root = self.root / "bundle_out"
        out_root.mkdir()
        with mock.patch.object(qt_app.QMessageBox, "information") as info, \
             mock.patch.object(qt_app.QFileDialog, "getExistingDirectory", return_value=str(out_root)):
            self.win.act_save_report_as.trigger()
        info.assert_not_called()
        copied = {p.name for p in (out_root / f"mcsa_{self.win.run_id}").iterdir()}
        expected = {
            Path(path).name
            for path in self.win.report_paths.values()
            if path and Path(path).exists()
        }
        self.assertEqual(copied, expected)

    def test_detail_json_still_loads_as_overlay(self):
        facts, paths = self._golden()
        payload = load_overlay_payload(Path(paths["evidence_json"]))
        events = extract_overlay_events(payload)
        self.assertEqual(len([e for e in events if e.get("accepted")]), len(facts["model"]["click_data"]))
        evidence = json.loads(Path(paths["evidence_json"]).read_text(encoding="utf-8"))
        self.assertEqual(evidence["run_status"], "COMPLETE")
        self.assertNotIn("metrics", evidence)
        self.assertNotIn("legacy_heuristic_assessment", evidence)
        self.assertNotIn("legacy_payload", evidence)

    def test_frozen_run_settings_and_events_not_mutated(self):
        self._run(_golden_offsets())
        locked = self.win.locked
        locked_before = copy.deepcopy(locked)
        events_before = copy.deepcopy(self.win.events)
        snapshot = self.win.host_evidence_snapshot
        snapshot_before = copy.deepcopy(snapshot)
        facts_before = self.win._completed_run_facts()
        self._generate()
        self.assertIs(self.win.locked, locked)
        self.assertEqual(self.win.locked, locked_before)
        self.assertEqual(self.win.events, events_before)
        self.assertIs(self.win.host_evidence_snapshot, snapshot)
        self.assertEqual(self.win.host_evidence_snapshot, snapshot_before)
        facts_after = self.win._completed_run_facts()
        self.assertEqual(facts_after["model"]["observations"], facts_before["model"]["observations"])
        self.assertEqual(facts_after["tags"], facts_before["tags"])

    def test_reports_use_frozen_model_after_setup_widgets_change(self):
        self._run(_golden_offsets())
        expected = self.win._completed_run_facts()["model"]["run"]["run_status"]
        self.win.device_edit.setText("OTHER_DUT")
        self.win.hz_combo.setCurrentText("40")
        self.win.timing_tol_edit.setText("9")
        self.win.radio_manual.setChecked(True)
        html_text, xlsx = self._artifacts(self._generate())
        self.assertEqual(_html_run_status(html_text)["Run Status"], expected)
        self.assertEqual(_xlsx_run_status(xlsx), expected)
        info = _html_kv(html_text, "test-information")
        self.assertEqual(info["Reference Rate"], "5.6 Hz")
        self.assertEqual(info["Device"], "123")
        self.assertNotIn("OTHER_DUT", html_text)


class ReportV2ModelTests(unittest.TestCase):
    def _settings(self, **extra):
        settings = legacy_app.normalize_settings({})
        settings.update({"profile_name": "Fixture_Custom", "fixture_hz": 10.0, "fixture_reference_hz": 10.0,
                         "duration_sec": 5, "focus_button": "Left", "expected_clicks": 0, "target_clicks": 0})
        settings.update(extra)
        return settings

    def _events(self, intervals):
        events, t = [], 0.0
        for i, step in enumerate(intervals):
            t += step if i else 0.0
            events.append({"index": i + 1, "button": "Left", "accepted": True, "event_type": "click",
                           "press_elapsed_ms": t, "release_elapsed_ms": t + 40.0, "elapsed_ms": t + 40.0,
                           "press_interval_ms": step if i else None, "hold_ms": 40.0})
        return events

    def test_report_code_filter_drops_non_finding_codes(self):
        self.assertEqual(report_v2.report_codes(["MANUAL_PROFILE_NO_EXPECTED_COUNT"]), [])
        self.assertEqual(
            report_v2.report_codes(["MANUAL_PROFILE_NO_EXPECTED_COUNT", "INTERVAL_JITTER_WARN"]),
            ["Interval variation"],
        )
        self.assertFalse(hasattr(report_v2, "report_result_codes"))

    def test_report_codes_use_neutral_names_without_physical_claims(self):
        codes = report_v2.report_codes
        self.assertEqual(codes(["MISS_OR_SLOW", "LONG_INTERVAL_SPIKE_WARN"]), ["Long interval"])
        self.assertEqual(codes(["LONG_INTERVAL_SPIKE_WARN", "MISS_OR_SLOW"]), ["Long interval"])
        self.assertEqual(codes(["MISS_OR_SLOW"]), ["Long interval"])
        self.assertEqual(codes(["HIGH_PROBABILITY_DOUBLE_CLICK", "TRUE_EXTRA_CLICK"]), ["Short interval"])
        self.assertEqual(codes(["INCOMPLETE_REBOUND", "FIXTURE_ACTUATION_TOO_SHORT"]), ["Short hold"])
        self.assertEqual(codes(["SETUP_ALIGNMENT_SUSPECT"]), ["Hold variation"])
        self.assertEqual(codes(["FIXTURE_HZ_MISMATCH"]), ["Cadence vs reference"])
        self.assertEqual(codes(["CHATTER_HARD", "TIMING_JITTER_PAIR"]), ["Very short interval", "Timing variation"])
        for tag in _PHYSICAL_CLAIM_CODES:
            self.assertNotIn(tag, report_v2.report_codes([tag, "LONG_INTERVAL_SPIKE_WARN"]), tag)
        self.assertEqual(report_v2.report_state("MISS_OR_SLOW"), "Long interval")
        self.assertEqual(report_v2.report_state("TRUE_EXTRA_CLICK"), "Short interval")
        self.assertEqual(report_v2.report_state("ON_CADENCE"), "Observation")

    def test_long_interval_run_is_complete_and_keeps_canonical_tag(self):
        intervals = [100.0] * 40
        intervals[20] = 200.0
        model = build_completed_result_model(locked=self._settings(), events=self._events(intervals),
                                             run_id="long_gap", target="U")
        self.assertIn("MISS_OR_SLOW", model["assessment"]["tags"])
        self.assertIn("LONG_INTERVAL_SPIKE_WARN", model["assessment"]["tags"])
        self.assertEqual(model["run"]["run_status"], "COMPLETE")
        before = copy.deepcopy(model)
        with tempfile.TemporaryDirectory() as tmp:
            html_path, xlsx_path = Path(tmp) / "r.html", Path(tmp) / "r.xlsx"
            report_v2.write_report_html(html_path, model)
            self.assertTrue(report_v2.write_report_xlsx(xlsx_path, model))
            html_text = html_path.read_text(encoding="utf-8")
            self.assertEqual(_html_run_status(html_text), {"Run Status": "COMPLETE"})
            self.assertEqual(_xlsx_run_status(xlsx_path), "COMPLETE")
            workbook = _workbook_text(xlsx_path)
            for text in (html_text, workbook):
                for tag in _PHYSICAL_CLAIM_CODES:
                    self.assertNotIn(tag, text, tag)
            wb = load_workbook(xlsx_path)
            try:
                headers = [c.value for c in wb["Click Data"][1]]
                click_text = " ".join(
                    "" if v is None else str(v)
                    for row in wb["Click Data"].iter_rows(min_row=2, values_only=True)
                    for v in row
                )
            finally:
                wb.close()
            self.assertNotIn("Observation", headers)
            self.assertNotIn("Long interval", click_text)
            self.assertNotIn("LONG_INTERVAL", click_text)
            self.assertNotIn("MISS_OR_SLOW", click_text)
            self.assertTrue(any("MISS_OR_SLOW" in (row.get("tags") or []) for row in model["click_data"]))
        self.assertEqual(model, before)
        self.assertIn("MISS_OR_SLOW", model["assessment"]["tags"])

    def test_available_device_identity_fields_appear_in_test_information(self):
        device = {"name": "HP Mouse 9000", "vid": "03F0", "pid": "0A4B", "bcd_device": "0102", "instance_id": "X"}
        model = build_completed_result_model(locked=self._settings(device_under_test=device),
                                             events=self._events([100.0] * 25), run_id="ident", target="SAMPLE-07")
        info = {label: value for _id, label, value in report_v2.build_report_view(model)["test_information"]}
        self.assertEqual(info["Run Profile"], "Fixture_Custom")
        self.assertEqual(info["Analysis scope"], "Measurement only")
        self.assertEqual(info["Device"], "SAMPLE-07")
        self.assertEqual(info["HID device name"], "HP Mouse 9000")
        self.assertEqual((info["VID"], info["PID"], info["bcdDevice"]), ("03F0", "0A4B", "0102"))
        self.assertNotIn("X", info.values())
        with tempfile.TemporaryDirectory() as tmp:
            xlsx_path = Path(tmp) / "r.xlsx"
            report_v2.write_report_xlsx(xlsx_path, model)
            summary = _sheet_kv(xlsx_path, "Summary")
        self.assertEqual((summary["VID"], summary["PID"], summary["bcdDevice"]), ("03F0", "0A4B", "0102"))
        host = {label: value for _id, label, value in report_v2.host_information_rows(
            {"platform": "P", "capture_backend": "C", "device_under_test": device, "schema": "MCSA_HOST_INFORMATION_V1"},
            tool_version="1",
        )}
        self.assertIn("Capture backend", host)
        self.assertIn("Tool version", host)
        self.assertIn("USB host controllers", host)
        self.assertNotIn("Host OS", host)

    def test_view_does_not_mutate_model_and_keeps_observations_in_model(self):
        intervals = [100.0] * 40
        intervals[25] = 15.0
        model = build_completed_result_model(locked=self._settings(), events=self._events(intervals),
                                             run_id="model_unit", target="U")
        self.assertTrue(model["observations"])
        before = copy.deepcopy(model)
        view = report_v2.build_report_view(model)
        self.assertEqual(model, before)
        self.assertNotIn("observations", view)
        self.assertNotIn("anomaly", view)
        html_text = report_v2.render_report_html(model, view)
        for forbidden in _EXPLANATORY:
            self.assertNotIn(forbidden, _visible_html_text(html_text), forbidden)

    def test_health_only_fail_cannot_render_silent_fail_report(self):
        from mouse_click_stability_analyzer.completed_result import qt_visible_verdict
        settings = self._settings()
        # Legacy heuristic diagnostic keeps its 4D-2a containment; it is no longer rendered.
        metrics = {
            "click_count": 20,
            "health": "FAIL",
            "fail_reasons": ["press_interval_outlier_ratio_pct 20.0000% > limit"],
        }
        self.assertEqual(qt_visible_verdict([], metrics, settings), "PASS")
        real_payload = legacy_app.build_payload

        def payload(*args, **kwargs):
            out = real_payload(*args, **kwargs)
            out["metrics"]["health"] = "FAIL"
            out["metrics"]["fail_reasons"] = list(metrics["fail_reasons"])
            return out

        events = self._events([100.0] * 25)
        with mock.patch("mouse_click_stability_analyzer.completed_result.build_payload", payload), \
             mock.patch("mouse_click_stability_analyzer.completed_result._metrics_issue_tags", return_value=[]):
            model = build_completed_result_model(
                locked=settings, events=events, run_id="health_only", target="U",
            )
        self.assertIsNotNone(model)
        self.assertEqual(model["run"]["run_status"], "COMPLETE")
        self.assertNotIn("verdict", model["assessment"])
        self.assertEqual(model["assessment"]["tags"], [])
        with tempfile.TemporaryDirectory() as tmp:
            html_path = Path(tmp) / "r.html"
            xlsx_path = Path(tmp) / "r.xlsx"
            report_v2.write_report_html(html_path, model)
            self.assertTrue(report_v2.write_report_xlsx(xlsx_path, model))
            html_text = html_path.read_text(encoding="utf-8")
            self.assertEqual(_html_run_status(html_text), {"Run Status": "COMPLETE"})
            self.assertEqual(_xlsx_run_status(xlsx_path), "COMPLETE")
            _assert_no_verdict_words(self, _visible_html_text(html_text))
            _assert_no_verdict_words(self, _workbook_text(xlsx_path))

    def test_fallback_without_model_uses_canonical_result(self):
        settings = self._settings(duration_sec=2)
        events = self._events([100.0] * 25)
        with tempfile.TemporaryDirectory() as tmp:
            paths = legacy_app.write_reports(Path(tmp), "fallback", "op", settings, events, "fallback_run")
            evidence = json.loads(Path(paths["evidence_json"]).read_text(encoding="utf-8"))
            self.assertNotIn("issue_tags", evidence)
            self.assertEqual(evidence["run_status"], "COMPLETE")
            html_text = Path(paths["engineering_html"]).read_text(encoding="utf-8")
            self.assertEqual(_html_run_status(html_text), {"Run Status": "COMPLETE"})
            for tag in _RETIRED:
                self.assertNotIn(tag, html_text)
            self.assertNotIn("EXCESS", _workbook_text(Path(paths["engineering_xlsx"])))
            self.assertTrue(legacy_app._verify_artifacts(paths)["ok"])


class ReportXlsxRowIdTests(unittest.TestCase):
    def test_xlsx_typing_is_keyed_by_row_id_not_visible_label(self):
        self.assertEqual(report_v2.xlsx_kv_cell_value("bcd_device", "0102"), "0102")
        self.assertIsInstance(report_v2.xlsx_kv_cell_value("bcd_device", "0102"), str)
        self.assertEqual(report_v2.xlsx_kv_cell_value("host_count_delta_vs_windows", "+2"), "+2")
        self.assertIsInstance(report_v2.xlsx_kv_cell_value("host_count_delta_vs_windows", "+2"), str)
        self.assertEqual(report_v2.xlsx_kv_cell_value("run_id", "20261006"), "20261006")
        self.assertEqual(report_v2.xlsx_kv_cell_value("host_press_count", "58"), 58)
        self.assertEqual(report_v2.xlsx_kv_cell_value("host_presses", "20"), 20)
        self.assertEqual(report_v2.xlsx_kv_cell_value("expected_slot_count", "56"), 56)

    def test_renamed_visible_label_does_not_change_xlsx_cell_type(self):
        device = {"name": "HP Mouse 9000", "vid": "03F0", "pid": "0A4B", "bcd_device": "0102"}
        settings = legacy_app.normalize_settings({})
        settings.update({
            "profile_name": "Fixture_Custom", "fixture_hz": 10.0, "fixture_reference_hz": 10.0,
            "duration_sec": 5, "focus_button": "Left", "expected_clicks": 0, "target_clicks": 0,
            "device_under_test": device,
        })
        events = []
        t = 0.0
        for i, step in enumerate([100.0] * 20):
            t += step if i else 0.0
            events.append({
                "index": i + 1, "button": "Left", "accepted": True, "event_type": "click",
                "press_elapsed_ms": t, "release_elapsed_ms": t + 40.0, "elapsed_ms": t + 40.0,
                "press_interval_ms": step if i else None, "hold_ms": 40.0,
            })
        model = build_completed_result_model(locked=settings, events=events, run_id="20261006", target="DUT")
        original = report_v2._test_information_rows

        def renamed(run, generated_at):
            return [
                (row_id, "Firmware rev" if row_id == "bcd_device" else label, value)
                for row_id, label, value in original(run, generated_at)
            ]

        with mock.patch.object(report_v2, "_test_information_rows", renamed):
            with tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "r.xlsx"
                self.assertTrue(report_v2.write_report_xlsx(path, model))
                summary = _sheet_kv(path, "Summary")
        self.assertEqual(summary["Firmware rev"], "0102")
        self.assertIsInstance(summary["Firmware rev"], str)
        self.assertNotIn("bcdDevice", summary)
        self.assertEqual(summary["Host Presses"], 20)


if __name__ == "__main__":
    unittest.main()
