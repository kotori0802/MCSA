"""Phase 4D-2h: Results / Report / Capture presentation containment."""
from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path

_SRC = Path(__file__).resolve().parents[1] / "src"
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from mouse_click_stability_analyzer.cadence import class_to_tags, classify_click
from mouse_click_stability_analyzer.completed_result import build_completed_result_model
from mouse_click_stability_analyzer.i18n import STRINGS, t
from mouse_click_stability_analyzer.qt_shell import app as qt_app
from mouse_click_stability_analyzer.presentation import (
    SCREENING_PRESENTATION_EN,
    SCREENING_PRESENTATION_KEYS,
    matched_observations_used_text,
    screening_presentation_label,
    screening_presentation_labels,
)
from mouse_click_stability_analyzer import report_v2
from mouse_click_stability_analyzer.host_evidence import build_fixture_host_evidence

_ALIAS_MAP = {
    "ON_CADENCE": "Observation",
    "CHATTER_BOUNCE": "Very short interval",
    "CHATTER_HARD": "Very short interval",
    "TRUE_EXTRA_CLICK": "Short interval",
    "HIGH_PROBABILITY_DOUBLE_CLICK": "Short interval",
    "INCOMPLETE_REBOUND": "Short hold",
    "FIXTURE_ACTUATION_TOO_SHORT": "Short hold",
    "LOW_HOLD_WARNING": "Short hold",
    "MISS_OR_SLOW": "Long interval",
    "LONG_INTERVAL_SPIKE_WARN": "Long interval",
}
_FORBIDDEN_HUMAN = (
    "Chatter detected",
    "Double-click detected",
    "Missed click",
    "DUT score",
    "Setup score",
    "Stability score",
    "Validation Profile",
    "Test Result PASS",
    "EXTRA OBSERVATION",
    "ISSUES",
)


def _make_qapp():
    try:
        from PySide6.QtWidgets import QApplication
    except Exception:  # pragma: no cover
        return None
    return QApplication.instance() or QApplication([])


def _tr(lang: str):
    return lambda key, **kwargs: t(lang, key, **kwargs)


class PresentationMappingTests(unittest.TestCase):
    def test_internal_enums_unchanged(self):
        self.assertEqual(classify_click(25.0, 80.0, {"fixture_hz": 40, "fixture_reference_hz": 40}), "ON_CADENCE")
        self.assertEqual(classify_click(5.0, 80.0, {"chatter_threshold_ms": 8}), "CHATTER_BOUNCE")
        self.assertEqual(
            classify_click(12.0, 80.0, {"chatter_threshold_ms": 8, "duplicate_click_threshold_ms": 30, "fixture_hz": 40}),
            "TRUE_EXTRA_CLICK",
        )
        self.assertEqual(class_to_tags("CHATTER_BOUNCE", False), ["CHATTER_HARD", "CHATTER_BOUNCE"])
        self.assertEqual(class_to_tags("TRUE_EXTRA_CLICK", False), ["HIGH_PROBABILITY_DOUBLE_CLICK"])
        self.assertEqual(class_to_tags("ON_CADENCE", False), [])

    def test_aliases_agree_across_surfaces(self):
        for internal, english in _ALIAS_MAP.items():
            with self.subTest(internal=internal):
                self.assertEqual(screening_presentation_label(internal), english)
                self.assertEqual(screening_presentation_label(internal, _tr("en")), english)
                self.assertEqual(report_v2.report_state(internal), english)
                self.assertEqual(SCREENING_PRESENTATION_EN[internal], english)
                self.assertIn(internal, SCREENING_PRESENTATION_KEYS)

    def test_deduped_flag_list(self):
        self.assertEqual(
            screening_presentation_labels(["TRUE_EXTRA_CLICK", "HIGH_PROBABILITY_DOUBLE_CLICK"], _tr("en")),
            ["Short interval"],
        )
        self.assertEqual(
            screening_presentation_labels(["CHATTER_HARD", "CHATTER_BOUNCE"], _tr("en")),
            ["Very short interval"],
        )

    def test_raw_enums_are_not_primary_labels(self):
        for internal in _ALIAS_MAP:
            label = screening_presentation_label(internal, _tr("en"))
            self.assertNotEqual(label, internal)
            self.assertNotIn("_", label)

    def test_zh_aliases_match_contract(self):
        self.assertEqual(screening_presentation_label("TRUE_EXTRA_CLICK", _tr("zh")), "短間隔")
        self.assertEqual(screening_presentation_label("CHATTER_HARD", _tr("zh")), "極短間隔")
        self.assertEqual(screening_presentation_label("INCOMPLETE_REBOUND", _tr("zh")), "短按住")
        self.assertEqual(screening_presentation_label("MISS_OR_SLOW", _tr("zh")), "長間隔")
        self.assertEqual(screening_presentation_label("ON_CADENCE", _tr("zh")), "觀測")
        self.assertEqual(STRINGS["zh"]["issues"], "觀測")
        self.assertEqual(STRINGS["en"]["issues"], "OBSERVATIONS")

    def test_matched_used_text_uses_in_capture_denominator(self):
        self.assertEqual(matched_observations_used_text(2, 55), "2 of 55 in capture")
        self.assertEqual(matched_observations_used_text(0, 58), "0 of 58 in capture")

    def test_report_layer_does_not_import_qt(self):
        import inspect
        from mouse_click_stability_analyzer import presentation as shared

        report_src = Path(inspect.getsourcefile(report_v2)).read_text(encoding="utf-8")
        shared_src = Path(inspect.getsourcefile(shared)).read_text(encoding="utf-8")
        self.assertNotIn("qt_shell", report_src)
        self.assertNotIn("PySide", report_src)
        self.assertNotIn("QApplication", report_src)
        self.assertNotIn("PySide", shared_src)
        self.assertNotIn("QtWidgets", shared_src)
        self.assertNotIn("QApplication", shared_src)


class ResultsReportSemanticsWindowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._qapp = _make_qapp()
        if cls._qapp is None:
            raise unittest.SkipTest("PySide6 QApplication unavailable")
        cls._tmp = tempfile.TemporaryDirectory()

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def setUp(self):
        self.win = qt_app.MainWindow(report_root=Path(self._tmp.name))
        self.win._listen_start_timer.stop()
        self.win.ui_timer.stop()
        self.win.ensure_raw_input = lambda: None
        self.win._cursor_in_test_area = lambda: True
        self.win._cursor_over_capture_exclusion_ui = lambda: False
        self.win.listening = True

    def tearDown(self):
        self.win.close()
        self.win.deleteLater()

    def _setup_fixture(self, *, hz="40", tol="2", auto_sec="1"):
        self.win.radio_fixture.setChecked(True)
        self.win.hz_combo.setCurrentText(hz)
        self.win.timing_tol_edit.setText(tol)
        self.win.auto_stop_on.setChecked(True)
        self.win.auto_stop_sec_combo.setCurrentText(auto_sec)

    def _run(self, offsets):
        self.assertTrue(self.win.lock_profile())
        self.assertTrue(self.win.arm_capture())
        t0 = float(self.win._arm_boundary_perf) + 0.001
        self.win._begin_session_from_first_click("Left", t0, t0 + 0.01, {})
        for off in offsets[1:]:
            press = t0 + off / 1000.0
            self.win.record_click("Left", press, press + 0.01, {})
        self.win.stop_run()
        if self.win._post_run_input_guard_active:
            self.win._release_post_run_input_guard()

    def _visible_text(self) -> str:
        parts = [
            self.win.issue_title.text(),
            self.win.results_context_label.text(),
            self.win.results_observations_text.text(),
            self.win.results_tags_label.text(),
            self.win.results_evidence_text.toPlainText(),
            self.win.results_timing_text.text(),
            self.win.results_verdict_label.text(),
        ]
        table = self.win.results_click_table
        for r in range(table.rowCount()):
            for c in range(table.columnCount()):
                item = table.item(r, c)
                if item:
                    parts.append(item.text())
        for c in range(table.columnCount()):
            header = table.horizontalHeaderItem(c)
            if header:
                parts.append(header.text())
        for key, val in self.win._results_evidence_cells:
            parts.append(key.text())
            parts.append(val.text())
        return "\n".join(parts)

    def test_click_data_and_capture_use_same_aliases(self):
        self._setup_fixture()
        self._run([0, 40, 80])
        facts = self.win._completed_run_facts()
        model_states = [row["state"] for row in facts["model"]["click_data"]]
        self.assertTrue(any(state in {"TRUE_EXTRA_CLICK", "ON_CADENCE"} for state in model_states))
        table = self.win.results_click_table
        displayed = [table.item(r, 5).text() for r in range(table.rowCount()) if table.item(r, 5)]
        for text in displayed:
            self.assertNotIn("TRUE_EXTRA_CLICK", text)
            self.assertNotIn("ON_CADENCE", text)
            self.assertNotIn("HIGH_PROBABILITY_DOUBLE_CLICK", text)
        self.assertTrue(any(text == "Short interval" or text == "Observation" for text in displayed))
        self.assertEqual(self.win._stat_name_labels["extra"].text(), "SHORT INTERVAL")
        self.assertEqual(self.win.issue_title.text(), "OBSERVATIONS")

    def test_interval_delta_is_not_matched_timing_error(self):
        self._setup_fixture(hz="40", tol="2")
        self._run([0, 40])
        facts = self.win._completed_run_facts()
        rows = facts["model"]["click_data"]
        self.assertAlmostEqual(float(rows[1]["interval_ms"]), 40.0, places=2)
        self.assertAlmostEqual(float(rows[1]["reference_delta_ms"]), 15.0, places=2)
        snap = facts["model"]["host_evidence"]["snapshot"]
        matched_n = int(snap["matching"]["matched_host_observation_count"])
        avg_err = snap["timing"]["matched_delta_stats"]["average"]
        self.assertGreaterEqual(matched_n, 1)
        self.assertNotAlmostEqual(float(rows[1]["reference_delta_ms"]), float(avg_err or 0), places=2)
        table = self.win.results_click_table
        self.assertEqual(table.horizontalHeaderItem(7).text(), "Interval Δ vs ref.")
        self.assertIn("Observed interval − configured reference interval", table.horizontalHeaderItem(7).toolTip())
        evidence = " ".join(f"{k.text()} {v.text()}" for k, v in self.win._results_evidence_cells)
        self.assertIn("Average matched timing error", evidence)
        self.assertNotIn("Ref. delta", evidence)
        self.assertNotIn("Error vs reference", self._visible_text())

    def test_matched_timing_denominator_zero_one_many_all(self):
        self._setup_fixture()
        self._run([0, 25, 50, 75])
        snap = dict(self.win.host_evidence_snapshot)
        matching = dict(snap["matching"])
        counts = dict(snap["counts"])
        none_stats = {"count": 0, "average": None, "minimum": None, "maximum": None, "stdev": None}

        matching["matched_host_observation_count"] = 0
        snap["matching"] = matching
        snap["timing"] = {"matched_delta_stats": dict(none_stats), "matched_absolute_delta_stats": dict(none_stats)}
        self.win.host_evidence_snapshot = snap
        self.win._refresh_results_workspace()
        evidence = " ".join(f"{k.text()} {v.text()}" for k, v in self.win._results_evidence_cells)
        total = counts["in_window_host_press_count"]
        self.assertIn("Matched observations used", evidence)
        self.assertIn(f"0 of {total} in capture", evidence)
        self.assertIn("No matched Host observations", evidence)
        self.assertNotIn("0.000 ms", evidence)

        matching["matched_host_observation_count"] = 1
        snap["matching"] = matching
        one = {"count": 1, "average": 1.5, "minimum": 1.5, "maximum": 1.5, "stdev": 0.0}
        snap["timing"] = {"matched_delta_stats": dict(one), "matched_absolute_delta_stats": dict(one)}
        self.win.host_evidence_snapshot = snap
        self.win._refresh_results_workspace()
        evidence = " ".join(f"{k.text()} {v.text()}" for k, v in self.win._results_evidence_cells)
        self.assertIn(f"1 of {total} in capture", evidence)
        self.assertIn("Statistics describe a single matched observation", evidence)
        self.assertIn("+1.500 ms", evidence)

        self.win.new_run()
        self._setup_fixture()
        self._run([0, 25, 50, 75])
        snap = self.win.host_evidence_snapshot
        matched = int(snap["matching"]["matched_host_observation_count"])
        total = int(snap["counts"]["in_window_host_press_count"])
        self.assertEqual(matched, total)
        evidence = " ".join(f"{k.text()} {v.text()}" for k, v in self.win._results_evidence_cells)
        self.assertIn(f"{matched} of {total} in capture", evidence)

    def test_representative_two_of_many_matched_population(self):
        self._setup_fixture(hz="5.6", auto_sec="10")
        offsets = [0.0]
        step = 174.0
        for _ in range(56):
            offsets.append(offsets[-1] + step)
        self._run(offsets)
        snap = self.win.host_evidence_snapshot
        self.assertIsNotNone(snap)
        matched = int(snap["matching"]["matched_host_observation_count"])
        total = int(snap["counts"]["in_window_host_press_count"])
        raw = int(snap["counts"]["raw_host_press_count"])
        self.assertGreaterEqual(total, 50)
        self.assertLess(matched, total)
        evidence = " ".join(f"{k.text()} {v.text()}" for k, v in self.win._results_evidence_cells)
        self.assertIn(f"{matched} of {total} in capture", evidence)
        self.assertIn("Matched observations used", evidence)
        timing = self.win.results_timing_text.text()
        self.assertIn(f"{matched} of {total} in capture", timing)
        view = report_v2.build_report_view(self.win._completed_run_facts()["model"])
        he = dict(item for group in view["host_evidence"] for item in group[1])
        self.assertEqual(he["Matched observations used"], f"{matched} of {total} in capture")
        if raw != total:
            self.assertNotIn(f"{matched} of {raw}", evidence)

    def test_matched_denominator_excludes_outside_capture(self):
        self._setup_fixture(auto_sec="1")
        self._run([0, 200, 400, 600, 1200])
        snap = self.win.host_evidence_snapshot
        self.assertIsNotNone(snap)
        raw = int(snap["counts"]["raw_host_press_count"])
        in_capture = int(snap["counts"]["in_window_host_press_count"])
        outside = int(snap["matching"]["outside_capture_window_observation_count"])
        matched = int(snap["matching"]["matched_host_observation_count"])
        unmatched = int(snap["matching"]["unmatched_host_observation_count"])
        ambiguous = int(snap["matching"]["ambiguous_host_observation_count"])
        self.assertEqual(raw, 5)
        self.assertEqual(outside, 1)
        self.assertEqual(in_capture, 4)
        self.assertEqual(matched + unmatched + ambiguous, in_capture)
        self.assertEqual(in_capture + outside, raw)
        evidence = " ".join(f"{k.text()} {v.text()}" for k, v in self.win._results_evidence_cells)
        used = f"{matched} of {in_capture} in capture"
        self.assertIn(used, evidence)
        self.assertNotIn(f"{matched} of {raw} in capture", evidence)
        timing = self.win.results_timing_text.text()
        self.assertIn(used, timing)
        view = report_v2.build_report_view(self.win._completed_run_facts()["model"])
        he = dict(item for group in view["host_evidence"] for item in group[1])
        self.assertEqual(he["Matched observations used"], used)
        self.assertEqual(he["Host observations in capture"], str(in_capture))
        self.assertEqual(he["Host observations recorded"], str(raw))
        self.assertEqual(he["Host observations outside capture"], str(outside))

    def test_legacy_scores_are_suppressed(self):
        self._setup_fixture()
        self._run([0, 25, 50, 75])
        facts = self.win._completed_run_facts()
        metrics = facts["metrics"]
        self.assertIn("dut_score", metrics)
        self.assertIn("setup_score", metrics)
        self.assertIsInstance(metrics["dut_score"], (int, float))
        self.assertIsInstance(metrics["setup_score"], (int, float))
        self.assertNotEqual(str(metrics["dut_score"]).lower(), "n/a")
        self.assertEqual(facts["model"]["assessment"]["metrics"]["dut_score"], metrics["dut_score"])
        visible = self._visible_text()
        self.assertNotIn("DUT 100", visible)
        self.assertNotIn("SETUP 100", visible)
        self.assertNotIn("Stability score", visible)
        self.assertNotIn("dut_score", visible)
        self.assertTrue(self.win.val_dut.parentWidget().isHidden())
        self.assertTrue(self.win.val_setup.parentWidget().isHidden())
        self.assertEqual(self.win.val_dut.text(), self.win._tr("na"))
        html, xlsx = self._report_text()
        for blob in (html, xlsx):
            self.assertNotIn("DUT score", blob)
            self.assertNotIn("SETUP score", blob)
            self.assertNotIn("Stability score", blob)
            self.assertNotIn("DUT 100", blob)

    def test_report_and_results_share_run_contract(self):
        self.win.device_edit.setText("SAMPLE")
        self._setup_fixture()
        self._run([0, 25, 50, 75])
        ctx = self.win.results_context_label.text()
        self.assertIn("Run Profile:", ctx)
        self.assertIn("Screening: Measurement only", ctx)
        self.assertEqual(self.win.results_verdict_label.text(), "COMPLETE")
        html, xlsx = self._report_text()
        for blob in (html, xlsx):
            self.assertIn("Run Status", blob)
            self.assertIn("Run Profile", blob)
            self.assertIn("Measurement only", blob)
            self.assertIn("Host-observed average interval", blob)
            self.assertIn("Average interval Δ vs reference", blob)
            self.assertNotIn("Validation Profile", blob)
            self.assertNotIn("Test Result", blob)
            self.assertNotIn("Chatter detected", blob)
            self.assertNotIn("Double-click detected", blob)
            self.assertNotIn("Missed click", blob)
            self.assertNotIn("Difference vs reference", blob)

    def test_observations_are_factual(self):
        self.win.radio_manual.setChecked(True)
        self.assertTrue(self.win.lock_profile())
        self.assertTrue(self.win.arm_capture())
        t0 = float(self.win._arm_boundary_perf) + 0.001
        self.win._begin_session_from_first_click("Left", t0, t0 + 0.003, {})
        self.win.record_click("Left", t0 + 0.005, t0 + 0.008, {})
        self.win.stop_run()
        if self.win._post_run_input_guard_active:
            self.win._release_post_run_input_guard()
        text = self.win.results_observations_text.text()
        self.assertNotIn("existing chatter rule", text)
        self.assertNotIn("existing double-click rule", text)
        self.assertNotIn("CHATTER_HARD", text)
        self.assertNotIn("HIGH_PROBABILITY_DOUBLE_CLICK", text)
        self.assertNotIn("Chatter detected", text)
        if text != self.win._tr("observations_none"):
            self.assertIn(self.win._tr("informational_observation"), text)

    def _report_text(self) -> tuple[str, str]:
        model = self.win._completed_run_facts()["model"]
        html_path = Path(self._tmp.name) / "sem.html"
        xlsx_path = Path(self._tmp.name) / "sem.xlsx"
        report_v2.write_report_html(html_path, model)
        report_v2.write_report_xlsx(xlsx_path, model)
        from openpyxl import load_workbook
        wb = load_workbook(xlsx_path)
        try:
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    parts.extend(str(v) for v in row if v is not None)
            xlsx = "\n".join(parts)
        finally:
            wb.close()
        return html_path.read_text(encoding="utf-8"), xlsx


class ReportViewSemanticsTests(unittest.TestCase):
    def test_click_columns_are_observation_oriented(self):
        cols = report_v2.click_data_columns(False)
        self.assertIn("Observation", cols)
        self.assertIn("Observation flag(s)", cols)
        self.assertIn("Interval Δ vs ref. (ms)", cols)
        self.assertNotIn("State", cols)
        self.assertNotIn("Rule Code(s)", cols)
        self.assertNotIn("Δ vs reference (ms)", cols)

    def test_zero_matched_omits_zero_quality_metrics(self):
        events = []
        t = 0.0
        for i, step in enumerate([0.0] + [40.0] * 8):
            t += step if i else 0.0
            events.append({
                "index": i + 1, "button": "Left", "accepted": True, "event_type": "click",
                "press_elapsed_ms": t, "release_elapsed_ms": t + 8.0, "elapsed_ms": t + 8.0,
                "press_interval_ms": step if i else None, "hold_ms": 8.0,
            })
        times = [e["press_elapsed_ms"] for e in events]
        snap = build_fixture_host_evidence(1.0, 40.0, 2.0, times)
        matching = dict(snap["matching"])
        matching["matched_host_observation_count"] = 0
        snap = dict(snap)
        snap["matching"] = matching
        snap["timing"] = {
            "matched_delta_stats": {"count": 0, "average": None, "minimum": None, "maximum": None, "stdev": None},
            "matched_absolute_delta_stats": {"count": 0, "average": None, "minimum": None, "maximum": None, "stdev": None},
        }
        locked = {
            "manual_mode": False, "profile_name": "Fixture_Custom", "focus_button": "Left",
            "fixture_hz": 40, "fixture_reference_hz": 40, "duration_sec": 1.0,
            "fixture_timing_tolerance_ms": 2.0, "button": "Left",
        }
        model = build_completed_result_model(
            locked=locked, events=events, run_id="zero", target="U", host_evidence_snapshot=snap,
        )
        groups = dict(report_v2._host_evidence_groups(model))
        timing = dict(groups["Matched timing (Host − schedule)"])
        self.assertTrue(timing["Matched observations used"].startswith("0 of "))
        self.assertEqual(timing["Matched timing statistics"], "No matched Host observations")
        self.assertNotIn("Average matched timing error", timing)


if __name__ == "__main__":
    unittest.main()
