"""Canonical Report V2: engineering HTML and XLSX rendered from the completed-run model.

The report is a factual test record: test configuration, host identity, the
run status, measured data, charts, and (XLSX) the recorded click table. No
DUT PASS / WARN / FAIL is rendered: no approved acceptance criterion exists.
Definitions, rule meanings, and interpretation guidance are User Guide content
and are never rendered here. Legacy expected-count inference, ordinal drift,
and artifact paths/filenames are never rendered either.
"""
from __future__ import annotations

import html
import math
import os
import re
from pathlib import Path
from typing import Any

from mouse_click_stability_analyzer.duration import (
    EXCEL_ELAPSED_NUMBER_FORMAT,
    elapsed_axis_uses_seconds,
    excel_elapsed_days,
    format_elapsed_axis_tick,
    format_elapsed_seconds_hms,
)
from mouse_click_stability_analyzer.chart_domain import (
    DOMAIN_HOLD,
    DOMAIN_INTERVAL,
    domain_axis_bounds,
)
from mouse_click_stability_analyzer.host_inventory import NOT_AVAILABLE, field_display
from mouse_click_stability_analyzer.schema import TOOL_VERSION
from mouse_click_stability_analyzer.presentation import (
    matched_observations_used_text,
    screening_presentation_label,
    screening_presentation_labels,
)
from mouse_click_stability_analyzer.interval_comparison import (
    POSITION_ABOVE,
    POSITION_BELOW,
    POSITION_WITHIN,
    classify_interval,
    summarize_positions,
)
from mouse_click_stability_analyzer.strings.loader import report_text as _rt

REPORT_TITLE = _rt("report_title")
_CHART_POINT_LIMIT = 1500
_NONE = _rt("report_none")

# Canonical tags that record a run property rather than a WARN/FAIL finding.
_NON_FINDING_RESULT_CODES = frozenset({"MANUAL_PROFILE_NO_EXPECTED_COUNT"})

# A canonical code dropped from the report when the listed code already records the same finding.
_REPORT_CODE_SUBSUMED_BY = {"MISS_OR_SLOW": "LONG_INTERVAL_SPIKE_WARN"}

_RUN_STATUS_FILLS = {"COMPLETE": "DDEBF7", "NO_DATA": "D9D9D9", "INVALID": "FCE4D6"}
# Rows whose values stay Excel text even when they look numeric (IDs, hex, signed counts).
_XLSX_TEXT_ROW_IDS = frozenset({
    "run_id", "run_profile", "device", "hid_device_name", "operator",
    "vid", "pid", "bcd_device", "host_count_delta_vs_windows",
    "os_version", "bios_version", "bios_date", "usb_host_controllers",
    "host_manufacturer", "host_model", "cpu_model", "python_runtime",
    "pyside6_version", "qt_version", "os_name",
})


def _finite(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _fmt(value: Any, digits: int = 2, unit: str = "", *, signed: bool = False) -> str:
    number = _finite(value)
    if number is None:
        return _NONE
    text = f"{number:+.{digits}f}" if signed and number != 0 else f"{number:.{digits}f}"
    return f"{text} {unit}" if unit else text


def _fmt_signed_count(value: Any) -> str:
    number = int(value or 0)
    return "0" if number == 0 else f"{number:+d}"


def xlsx_kv_cell_value(row_id: str, value: Any) -> Any:
    """Excel cell value for a kv row. Typing is keyed by stable row_id, never by visible label."""
    if row_id not in _XLSX_TEXT_ROW_IDS and isinstance(value, str) and re.fullmatch(r"[+-]?\d+", value):
        return int(value)
    return value


def report_codes(tags: list[str]) -> list[str]:
    """Report-presentation names for canonical codes: neutral aliases, subsumed codes dropped, deduplicated."""
    present = {str(t) for t in tags}
    kept: list[str] = []
    for tag in tags:
        name = str(tag)
        if name in _NON_FINDING_RESULT_CODES or _REPORT_CODE_SUBSUMED_BY.get(name) in present:
            continue
        kept.append(name)
    return screening_presentation_labels(kept)


def report_state(state: str) -> str:
    return screening_presentation_label(state)


def report_has_host_evidence(model: dict[str, Any]) -> bool:
    return not model["run"].get("manual") and bool((model.get("host_evidence") or {}).get("snapshot"))


def report_sheet_names(model: dict[str, Any]) -> tuple[str, ...]:
    if report_has_host_evidence(model):
        return (
            _rt("report_sheet_summary"),
            _rt("report_sheet_measurements"),
            _rt("report_sheet_host_evidence"),
            _rt("report_sheet_click_data"),
        )
    return (
        _rt("report_sheet_summary"),
        _rt("report_sheet_measurements"),
        _rt("report_sheet_click_data"),
    )


def _run_status_rows(run_status: str) -> list[tuple[str, str, str]]:
    return [("run_status", _rt("report_label_run_status"), run_status)]


def _present(rows: list[tuple[str, str, Any]]) -> list[tuple[str, str, str]]:
    out = []
    for row_id, label, value in rows:
        text = str(value if value is not None else "").strip()
        if text and text != _NONE:
            out.append((row_id, label, text))
    return out


def _test_information_rows(run: dict[str, Any], generated_at: str) -> list[tuple[str, str, str]]:
    settings = run.get("settings") or {}
    device = settings.get("device_under_test") or {}
    rows: list[tuple[str, str, Any]] = [
        ("run_id", _rt("report_label_run_id"), run.get("run_id")),
        ("run_profile", _rt("report_label_run_profile"), settings.get("profile_name")),
        ("screening", _rt("report_label_screening"), _rt("report_value_measurement_only")),
        ("device", _rt("report_label_device"), run.get("target")),
        ("hid_device_name", _rt("report_label_hid_device_name"), device.get("name")),
        ("vid", _rt("report_label_vid"), device.get("vid")),
        ("pid", _rt("report_label_pid"), device.get("pid")),
        ("bcd_device", _rt("report_label_bcd_device"), device.get("bcd_device")),
        ("button", _rt("report_label_button"), run.get("button")),
        ("mode", _rt("report_label_mode"), _rt("report_value_manual") if run.get("manual") else _rt("report_value_fixture")),
    ]
    if not run.get("manual"):
        hz = _finite(run.get("reference_hz"))
        tol = _finite(run.get("timing_tolerance_ms"))
        ref = _finite(run.get("reference_interval_ms"))
        relative = tol / ref * 100.0 if tol is not None and ref else None
        rows += [
            ("reference_rate", _rt("report_label_reference_rate"), f"{hz:g} Hz" if hz is not None else None),
            ("reference_interval", _rt("report_label_reference_interval"), _fmt(ref, 2, "ms")),
            ("schedule_association_window", _rt("report_label_schedule_association_window"), f"±{tol:.2f} ms" if tol is not None else None),
            ("association_relative", _rt("report_label_association_relative"), f"{relative:.1f} %" if relative is not None else None),
        ]
        rows.extend(_interval_comparison_info_rows(run.get("interval_comparison")))
    elapsed = _finite(run.get("elapsed_s"))
    rows += [
        ("duration", _rt("report_label_duration"), format_elapsed_seconds_hms(elapsed) if elapsed is not None else None),
        ("host_presses", _rt("report_label_host_presses"), str(int(run.get("host_press_count") or 0))),
        ("operator", _rt("report_label_operator"), run.get("operator")),
        ("report_generated", _rt("report_label_report_generated"), generated_at),
    ]
    return _present(rows)


_HOST_INFO_GROUPS: tuple[tuple[str, tuple[tuple[str, str], ...]], ...] = (
    ("report_group_host_platform", (
        ("host_manufacturer", "report_label_host_manufacturer"),
        ("host_model", "report_label_host_model"),
    )),
    ("report_group_host_os", (
        ("os_name", "report_label_os_name"),
        ("os_version", "report_label_os_version"),
        ("cpu_model", "report_label_cpu_model"),
    )),
    ("report_group_host_firmware", (
        ("bios_vendor", "report_label_bios_vendor"),
        ("bios_version", "report_label_bios_version"),
        ("bios_date", "report_label_bios_date"),
    )),
    ("report_group_host_usb", (
        ("usb_host_controllers", "report_label_usb_host_controllers"),
        ("dut_usb_controller_attributed", "report_label_dut_usb_attribution"),
        ("capture_backend", "report_label_capture_backend"),
    )),
    ("report_group_host_tool_runtime", (
        ("tool_version", "report_label_tool_version"),
        ("python_runtime", "report_label_python_runtime"),
        ("pyside6_version", "report_label_pyside6_version"),
        ("qt_version", "report_label_qt_version"),
    )),
)


def host_information_groups(host_env: dict[str, Any] | None, *, tool_version: str) -> list[tuple[str, list[tuple[str, str, str]]]]:
    env = dict(host_env or {})
    if tool_version and not env.get("tool_version"):
        env["tool_version"] = tool_version
    groups: list[tuple[str, list[tuple[str, str, str]]]] = []
    for group_key, fields in _HOST_INFO_GROUPS:
        rows: list[tuple[str, str, str]] = []
        for field_id, label_key in fields:
            if field_id == "dut_usb_controller_attributed":
                attributed = bool(env.get("dut_usb_controller_attributed"))
                value = _rt("report_value_dut_usb_attributed") if attributed else _rt("report_value_dut_usb_not_attributed")
            elif field_id == "tool_version":
                value = field_display(env, "tool_version") if env.get("tool_version") else str(tool_version or "")
            else:
                value = field_display(env, field_id)
            if value == NOT_AVAILABLE:
                value = _rt("report_value_not_available")
            rows.append((field_id, _rt(label_key), value))
        groups.append((_rt(group_key), rows))
    return groups


def host_information_rows(host_env: dict[str, Any] | None, *, tool_version: str) -> list[tuple[str, str, str]]:
    rows: list[tuple[str, str, str]] = []
    for _group, group_rows in host_information_groups(host_env, tool_version=tool_version):
        rows.extend(group_rows)
    return rows


def _measurement_rows(model: dict[str, Any]) -> list[tuple[str, str, str]]:
    run = model["run"]
    metrics = model["assessment"]["metrics"]
    avg = _finite(metrics.get("interval_avg_ms"))
    rows: list[tuple[str, str, str]] = [
        ("host_average_interval", _rt("report_label_host_average_interval"), _fmt(avg, 2, "ms")),
        ("host_cadence", _rt("report_label_host_cadence"), _fmt(1000.0 / avg if avg else None, 2, "Hz")),
    ]
    if not run.get("manual"):
        ref = _finite(run.get("reference_interval_ms"))
        rows.append(("measurement_reference_interval", _rt("report_label_measurement_reference_interval"), _fmt(ref, 2, "ms")))
        diff = avg - ref if avg is not None and ref is not None else None
        rows.append(("average_interval_delta", _rt("report_label_average_interval_delta"), _fmt(diff, 2, "ms", signed=True)))
    cv = _finite(metrics.get("interval_cv_pct"))
    if cv is not None:
        rows.append(("interval_cv", _rt("report_label_interval_cv"), _fmt(cv, 2, "%")))
    rows.append(("average_hold", _rt("report_label_average_hold"), _fmt(metrics.get("hold_avg_ms"), 2, "ms")))
    rows.append(("host_press_count", _rt("report_label_host_press_count"), str(len(model["click_data"]))))
    rows.extend(_interval_comparison_measurement_rows(model))
    return rows


def _comparison_enabled(contract: dict[str, Any] | None) -> bool:
    return bool(contract and contract.get("enabled"))


def _interval_comparison_info_rows(contract: dict[str, Any] | None) -> list[tuple[str, str, Any]]:
    if not _comparison_enabled(contract):
        return []
    basis = str(contract.get("basis") or "")
    configured = _finite(contract.get("configured_value"))
    if basis == "percent":
        tolerance = _rt("report_fmt_tolerance_pct", value=f"{configured:.2f}") if configured is not None else None
        derived = _rt("report_fmt_tolerance_ms", value=f"{float(contract['tolerance_ms']):.2f}")
        basis_label = _rt("report_value_basis_percent")
    else:
        tolerance = _rt("report_fmt_tolerance_ms", value=f"{configured:.2f}") if configured is not None else None
        derived = _rt("report_fmt_tolerance_pct", value=f"{float(contract['tolerance_pct']):.2f}")
        basis_label = _rt("report_value_basis_milliseconds")
    bounds = _rt(
        "report_fmt_interval_range",
        lo=f"{float(contract['lower_bound_ms']):.2f}",
        hi=f"{float(contract['upper_bound_ms']):.2f}",
    )
    return [
        ("interval_comparison", _rt("report_label_interval_comparison"), _rt("report_value_interval_comparison_enabled")),
        ("range_basis", _rt("report_label_range_basis"), basis_label),
        ("configured_tolerance", _rt("report_label_configured_tolerance"), tolerance),
        ("derived_equivalent", _rt("report_label_derived_equivalent"), derived),
        ("configured_interval_range", _rt("report_label_configured_interval_range"), bounds),
    ]


def _interval_comparison_measurement_rows(model: dict[str, Any]) -> list[tuple[str, str, str]]:
    contract = model.get("interval_comparison") or (model.get("run") or {}).get("interval_comparison")
    if not _comparison_enabled(contract):
        return []
    positions = [str(row.get("range_position") or "") for row in model.get("click_data") or []]
    summary = summarize_positions(positions)
    within_pct = summary.get("within_pct")
    pct_text = f"{within_pct:.2f} %" if within_pct is not None else _NONE
    return [
        ("comparable_intervals", _rt("report_label_comparable_intervals"), str(int(summary["comparable_interval_count"] or 0))),
        ("within_configured_range", _rt("report_label_within_configured_range"), str(int(summary["within_count"] or 0))),
        ("outside_configured_range", _rt("report_label_outside_configured_range"), str(int(summary["outside_count"] or 0))),
        ("below_configured_range", _rt("report_label_below_configured_range"), str(int(summary["below_count"] or 0))),
        ("above_configured_range", _rt("report_label_above_configured_range"), str(int(summary["above_count"] or 0))),
        ("within_configured_range_pct", _rt("report_label_within_configured_range_pct"), pct_text),
    ]


def _range_position_text(token: str) -> str:
    mapping = {
        POSITION_WITHIN: "report_value_range_within",
        POSITION_BELOW: "report_value_range_below",
        POSITION_ABOVE: "report_value_range_above",
    }
    key = mapping.get(str(token or ""))
    return _rt(key) if key else _rt("report_value_range_blank")


def _host_evidence_groups(model: dict[str, Any]) -> list[tuple[str, list[tuple[str, str, str]]]]:
    if not report_has_host_evidence(model):
        return []
    snapshot = model["host_evidence"]["snapshot"]
    config = snapshot["configuration"]
    counts = snapshot["counts"]
    matching = snapshot["matching"]
    delta = snapshot["timing"]["matched_delta_stats"]
    abs_delta = snapshot["timing"]["matched_absolute_delta_stats"]
    matched_n = int(matching["matched_host_observation_count"] or 0)
    in_capture_n = int(counts["in_window_host_press_count"] or 0)
    timing_rows: list[tuple[str, str, str]] = [
        ("matched_observations_used", _rt("report_label_matched_observations_used"), matched_observations_used_text(matched_n, in_capture_n)),
    ]
    if matched_n == 0:
        timing_rows.append(("matched_timing_empty", _rt("report_label_matched_timing_statistics"), _rt("report_value_no_matched_host_observations")))
    else:
        timing_rows.extend([
            ("average_matched_timing_error", _rt("report_label_average_matched_timing_error"), _fmt(delta["average"], 3, "ms", signed=True)),
            ("average_abs_matched_timing_error", _rt("report_label_average_abs_matched_timing_error"), _fmt(abs_delta["average"], 3, "ms")),
            ("min_matched_timing_error", _rt("report_label_min_matched_timing_error"), _fmt(delta["minimum"], 3, "ms", signed=True)),
            ("max_matched_timing_error", _rt("report_label_max_matched_timing_error"), _fmt(delta["maximum"], 3, "ms", signed=True)),
            ("matched_timing_stdev", _rt("report_label_matched_timing_stdev"), _fmt(delta["stdev"], 3, "ms")),
        ])
        if matched_n == 1:
            timing_rows.append(("matched_timing_single", _rt("report_label_note"), _rt("report_note_single_matched")))
    return [
        (_rt("report_group_schedule"), [
            ("expected_slot_count", _rt("report_label_schedule_windows"), str(counts["expected_slot_count"])),
            ("configured_reference_interval", _rt("report_label_configured_reference_interval"), _fmt(config["reference_interval_ms"], 2, "ms")),
            ("schedule_association_window", _rt("report_label_schedule_association_window"), f"±{float(config['timing_tolerance_ms']):.3f} ms"),
        ]),
        (_rt("report_group_host_observations"), [
            ("in_capture", _rt("report_label_in_capture"), str(counts["in_window_host_press_count"])),
            ("recorded", _rt("report_label_recorded"), str(counts["raw_host_press_count"])),
            ("host_count_delta_vs_windows", _rt("report_label_host_count_delta_vs_windows"), _fmt_signed_count(counts["host_count_delta_vs_expected_slots"])),
        ]),
        (_rt("report_group_occupancy"), [
            ("unoccupied_windows", _rt("report_label_unoccupied_windows"), str(matching["empty_expected_slot_count"])),
            ("single_occupancy_windows", _rt("report_label_single_occupancy_windows"), str(matching["single_occupancy_slot_count"])),
            ("multi_occupancy_windows", _rt("report_label_multi_occupancy_windows"), str(matching["multi_occupancy_slot_count"])),
        ]),
        (_rt("report_group_association"), [
            ("matched_host_observations", _rt("report_label_matched_host_observations"), str(matching["matched_host_observation_count"])),
            ("unmatched_host_observations", _rt("report_label_unmatched_host_observations"), str(matching["unmatched_host_observation_count"])),
            ("ambiguous_host_observations", _rt("report_label_ambiguous_host_observations"), str(matching["ambiguous_host_observation_count"])),
            ("outside_capture", _rt("report_label_outside_capture"), str(matching["outside_capture_window_observation_count"])),
        ]),
        (_rt("report_group_matched_timing"), timing_rows),
    ]


def click_data_columns(
    manual: bool,
    *,
    include_classification: bool = False,
    include_range_position: bool = False,
) -> list[str]:
    cols = [
        _rt("report_click_col_index"),
        _rt("report_click_col_elapsed"),
        _rt("report_click_col_button"),
        _rt("report_click_col_interval"),
        _rt("report_click_col_hold"),
    ]
    if include_classification:
        cols.append(_rt("report_click_col_observation"))
        cols.append(_rt("report_click_col_flags"))
    if not manual:
        cols.append(_rt("report_click_col_ref_delta"))
    if include_range_position:
        cols.append(_rt("report_click_col_range_position"))
    return cols


def _click_values(
    row: dict[str, Any],
    manual: bool,
    *,
    include_classification: bool = False,
    include_range_position: bool = False,
) -> list[Any]:
    values: list[Any] = [
        row.get("index"),
        _finite(row.get("elapsed_ms")),
        row.get("button") or "",
        _finite(row.get("interval_ms")),
        _finite(row.get("hold_ms")),
    ]
    if include_classification:
        values.append(report_state(row.get("state") or ""))
        values.append(", ".join(report_codes(list(row.get("tags") or []))))
    if not manual:
        values.append(_finite(row.get("reference_delta_ms")))
    if include_range_position:
        values.append(_range_position_text(str(row.get("range_position") or "")))
    return values


def build_report_view(
    model: dict[str, Any],
    *,
    host_info: list[tuple[str, str, str]] | None = None,
    generated_at: str = "",
) -> dict[str, Any]:
    """Single presentation view shared by the HTML and XLSX writers. Does not mutate ``model``."""
    run = model["run"]
    manual = bool(run.get("manual"))
    run_status = str(run["run_status"])
    ref = None if manual else _finite(model["analysis"].get("reference_interval_ms"))
    comparison = model.get("interval_comparison") or run.get("interval_comparison")
    chart_reference = None
    if ref is not None:
        chart_reference = {"interval_ms": ref}
        if _comparison_enabled(comparison):
            chart_reference["lower_bound_ms"] = float(comparison["lower_bound_ms"])
            chart_reference["upper_bound_ms"] = float(comparison["upper_bound_ms"])
    include_range = (not manual) and _comparison_enabled(comparison)
    frozen_host = run.get("host_information") if isinstance(run.get("host_information"), dict) else None
    groups = host_information_groups(frozen_host, tool_version=TOOL_VERSION) if frozen_host or host_info else []
    rows = list(host_info) if host_info is not None else (
        host_information_rows(frozen_host, tool_version=TOOL_VERSION) if frozen_host else []
    )
    if not groups and rows:
        groups = [(_rt("report_section_host_information"), rows)]
    return {
        "chart_reference": chart_reference,
        "title": _rt("report_title"),
        "run_status": run_status,
        "run_status_rows": _run_status_rows(run_status),
        "test_information": _test_information_rows(run, generated_at),
        "host_information": rows,
        "host_information_groups": groups,
        "measurements": _measurement_rows(model),
        "host_evidence": _host_evidence_groups(model),
        "click_columns": click_data_columns(manual, include_range_position=include_range),
        "click_rows": [_click_values(r, manual, include_range_position=include_range) for r in model["click_data"]],
        "click_range_positions": [str(r.get("range_position") or "") for r in model["click_data"]] if include_range else [],
        "interval_observation_count": len(model["analysis"].get("intervals") or []),
        "manual": manual,
    }


# --------------------------------------------------------------------------- HTML

def _e(value: Any) -> str:
    return html.escape(str(value if value is not None else ""), quote=True)


def _kv_table(rows: list[tuple[str, str, str]]) -> str:
    body = "".join(
        f"<tr><th scope='row'>{_e(label)}</th><td>{_e(value)}</td></tr>"
        for _row_id, label, value in rows
    )
    return f"<table class='kv'>{body}</table>"


def _downsample(values: list[Any]) -> list[Any]:
    if len(values) <= _CHART_POINT_LIMIT:
        return list(values)
    step = len(values) / _CHART_POINT_LIMIT
    return [values[int(i * step)] for i in range(_CHART_POINT_LIMIT)]


_W, _H, _L, _R, _T, _B = 600, 240, 56, 12, 12, 34


def _axis(
    xmin: float,
    xmax: float,
    ymin: float,
    ymax: float,
    xlabel: str,
    ylabel: str,
    *,
    format_x: Any = None,
) -> str:
    x0, x1, y0, y1 = _L, _W - _R, _H - _B, _T
    xmin_text = format_x(xmin) if format_x else f"{xmin:.0f}"
    xmax_text = format_x(xmax) if format_x else f"{xmax:.0f}"
    return (
        f"<line x1='{x0}' y1='{y0}' x2='{x1}' y2='{y0}' class='ax'/>"
        f"<line x1='{x0}' y1='{y0}' x2='{x0}' y2='{y1}' class='ax'/>"
        f"<text x='{x0}' y='{y0 + 14}' class='tk'>{_e(xmin_text)}</text>"
        f"<text x='{x1}' y='{y0 + 14}' class='tk' text-anchor='end'>{_e(xmax_text)}</text>"
        f"<text x='{x0 - 4}' y='{y0}' class='tk' text-anchor='end'>{ymin:.1f}</text>"
        f"<text x='{x0 - 4}' y='{y1 + 8}' class='tk' text-anchor='end'>{ymax:.1f}</text>"
        f"<text x='{(x0 + x1) / 2:.0f}' y='{_H - 4}' class='lb' text-anchor='middle'>{_e(xlabel)}</text>"
        f"<text x='12' y='{(y0 + y1) / 2:.0f}' class='lb' text-anchor='middle' "
        f"transform='rotate(-90 12 {(y0 + y1) / 2:.0f})'>{_e(ylabel)}</text>"
    )


def _scale(v: float, lo: float, hi: float, a: float, b: float) -> float:
    return a if hi <= lo else a + (v - lo) / (hi - lo) * (b - a)


def _padded(lo: float, hi: float) -> tuple[float, float]:
    pad = (hi - lo) * 0.08 or max(1.0, abs(hi) * 0.05)
    return lo - pad, hi + pad


def _svg(body: str, title: str) -> str:
    return (
        f"<svg viewBox='0 0 {_W} {_H}' role='img' aria-label='{_e(title)}' "
        f"preserveAspectRatio='xMidYMid meet'>{body}</svg>"
    )


def _empty_chart(title: str) -> str:
    return _svg(f"<text x='{_W / 2:.0f}' y='{_H / 2:.0f}' class='lb' text-anchor='middle'>{_e(_rt('report_chart_no_data'))}</text>", title)


# Report charts are interval-domain: they draw the configured reference and,
# when enabled, the configured interval range. Schedule Association Window
# applies to timestamps and has no interval band.
def _configured_bounds(reference: dict[str, Any] | None) -> tuple[float, float] | None:
    if not reference:
        return None
    lo = _finite(reference.get("lower_bound_ms"))
    hi = _finite(reference.get("upper_bound_ms"))
    if lo is None or hi is None:
        return None
    return lo, hi


def _reference_span(reference: dict[str, Any] | None) -> tuple[float, float] | None:
    if not reference:
        return None
    ref = reference["interval_ms"]
    bounds = _configured_bounds(reference)
    if bounds:
        return min(bounds[0], ref), max(bounds[1], ref)
    return ref, ref


def _range_contract(reference: dict[str, Any] | None) -> dict[str, Any] | None:
    bounds = _configured_bounds(reference)
    if not bounds:
        return None
    return {"enabled": True, "lower_bound_ms": bounds[0], "upper_bound_ms": bounds[1]}


def _reference_label(reference: dict[str, Any]) -> str:
    return f"{reference['interval_ms']:.2f} ms"


def _svg_diamond(cx: float, cy: float, *, r: float = 3.2) -> str:
    return (
        f"<polygon class='out' points='{cx:.1f},{cy - r:.1f} {cx + r:.1f},{cy:.1f} "
        f"{cx:.1f},{cy + r:.1f} {cx - r:.1f},{cy:.1f}'/>"
    )


def _svg_range_band_horizontal(reference: dict[str, Any], ymin: float, ymax: float) -> str:
    bounds = _configured_bounds(reference)
    if not bounds:
        return ""
    lo, hi = bounds
    y0 = _scale(hi, ymin, ymax, _H - _B, _T)
    y1 = _scale(lo, ymin, ymax, _H - _B, _T)
    top, bot = (y0, y1) if y0 <= y1 else (y1, y0)
    return (
        f"<rect class='band' x='{_L}' y='{top:.1f}' width='{_W - _R - _L}' height='{max(0.0, bot - top):.1f}'/>"
        f"<line x1='{_L}' y1='{y0:.1f}' x2='{_W - _R}' y2='{y0:.1f}' class='bnd'/>"
        f"<line x1='{_L}' y1='{y1:.1f}' x2='{_W - _R}' y2='{y1:.1f}' class='bnd'/>"
    )


def _svg_range_band_vertical(reference: dict[str, Any], xmin: float, xmax: float) -> str:
    bounds = _configured_bounds(reference)
    if not bounds:
        return ""
    lo, hi = bounds
    x0 = _scale(lo, xmin, xmax, _L, _W - _R)
    x1 = _scale(hi, xmin, xmax, _L, _W - _R)
    left, right = (x0, x1) if x0 <= x1 else (x1, x0)
    return (
        f"<rect class='band' x='{left:.1f}' y='{_T}' width='{max(0.0, right - left):.1f}' height='{_H - _B - _T}'/>"
        f"<line x1='{x0:.1f}' y1='{_T}' x2='{x0:.1f}' y2='{_H - _B}' class='bnd'/>"
        f"<line x1='{x1:.1f}' y1='{_T}' x2='{x1:.1f}' y2='{_H - _B}' class='bnd'/>"
    )


def _svg_reference_horizontal(reference: dict[str, Any], ymin: float, ymax: float) -> str:
    ref = reference["interval_ms"]
    ry = _scale(ref, ymin, ymax, _H - _B, _T)
    return (f"<line x1='{_L}' y1='{ry:.1f}' x2='{_W - _R}' y2='{ry:.1f}' class='ref'>"
            f"<title>{_rt("report_legend_configured_reference")}</title></line>"
            f"<text x='{_W - _R - 4}' y='{ry - 5:.1f}' class='rl' text-anchor='end'>{_reference_label(reference)}</text>")


def _svg_reference_vertical(reference: dict[str, Any], xmin: float, xmax: float) -> str:
    ref = reference["interval_ms"]
    rx = _scale(ref, xmin, xmax, _L, _W - _R)
    anchor = "start" if rx < (_L + _W - _R) / 2 else "end"
    dx = 5 if anchor == "start" else -5
    return (f"<line x1='{rx:.1f}' y1='{_T}' x2='{rx:.1f}' y2='{_H - _B}' class='ref'>"
            f"<title>{_rt("report_legend_configured_reference")}</title></line>"
            f"<text x='{rx + dx:.1f}' y='{_T + 12}' class='rl' text-anchor='{anchor}'>{_reference_label(reference)}</text>")


def _svg_interval_series(series: list[tuple[float, float]], reference: dict[str, Any] | None) -> str:
    title = _rt("report_chart_interval_over_time")
    pts = _downsample(series)
    if len(pts) < 2:
        return _empty_chart(title)
    xs = [p[0] for p in pts]
    ys = [p[1] for p in pts]
    ymin, ymax = min(ys), max(ys)
    span = _reference_span(reference)
    if span:
        ymin, ymax = min(ymin, span[0]), max(ymax, span[1])
    ymin, ymax = domain_axis_bounds(*_padded(ymin, ymax), domain=DOMAIN_INTERVAL, values=ys)
    xmin, xmax = min(xs), max(xs)
    span_ms = max(xmax - xmin, abs(xmax), abs(xmin))
    x_label = _rt("report_axis_elapsed_s") if elapsed_axis_uses_seconds(span_ms) else _rt("report_axis_elapsed")
    coords = " ".join(
        f"{_scale(x, xmin, xmax, _L, _W - _R):.1f},{_scale(y, ymin, ymax, _H - _B, _T):.1f}"
        for x, y in zip(xs, ys)
    )
    body = _axis(
        xmin, xmax, ymin, ymax, x_label, _rt("report_axis_interval_ms"),
        format_x=lambda v, span=span_ms: format_elapsed_axis_tick(v, span_ms=span),
    )
    if reference:
        body += _svg_range_band_horizontal(reference, ymin, ymax)
        body += _svg_reference_horizontal(reference, ymin, ymax)
    body += f"<polyline points='{coords}' class='series'><title>{_rt("report_legend_host_observation")}</title></polyline>"
    contract = _range_contract(reference)
    if contract:
        for x, y in zip(xs, ys):
            if classify_interval(y, contract) in (POSITION_BELOW, POSITION_ABOVE):
                body += _svg_diamond(
                    _scale(x, xmin, xmax, _L, _W - _R),
                    _scale(y, ymin, ymax, _H - _B, _T),
                )
    return _svg(body, title)


def _svg_histogram(bins: list[dict[str, Any]], reference: dict[str, Any] | None) -> str:
    title = _rt("report_chart_interval_distribution")
    if not bins:
        return _empty_chart(title)
    lo, hi = float(bins[0]["lo"]), float(bins[-1]["hi"])
    span = _reference_span(reference)
    if span:
        lo, hi = min(lo, span[0]), max(hi, span[1])
    peak = max(int(b["count"]) for b in bins) or 1
    body = _axis(lo, hi, 0, peak, _rt("report_axis_interval_ms"), _rt("report_axis_count"))
    if reference:
        body += _svg_range_band_vertical(reference, lo, hi)
    for b in bins:
        h = (int(b["count"]) / peak) * (_H - _B - _T)
        x0 = _scale(float(b["lo"]), lo, hi, _L, _W - _R)
        x1 = _scale(float(b["hi"]), lo, hi, _L, _W - _R)
        body += (
            f"<rect x='{x0 + 1:.1f}' y='{_H - _B - h:.1f}' width='{max(1.0, x1 - x0 - 2):.1f}' "
            f"height='{h:.1f}' class='bar'><title>{float(b['lo']):.2f}–{float(b['hi']):.2f} ms: "
            f"{int(b['count'])}</title></rect>"
        )
    if reference:
        body += _svg_reference_vertical(reference, lo, hi)
    return _svg(body, title)


def _svg_scatter(points: list[tuple[float, float, str]], reference: dict[str, Any] | None) -> str:
    title = _rt("report_chart_hold_vs_interval")
    pts = _downsample(points)
    if not pts:
        return _empty_chart(title)
    xs = [p[0] for p in pts]
    ys = [p[1] for p in pts]
    xlo, xhi = min(xs), max(xs)
    span = _reference_span(reference)
    if span:
        xlo, xhi = min(xlo, span[0]), max(xhi, span[1])
    xmin, xmax = domain_axis_bounds(*_padded(xlo, xhi), domain=DOMAIN_INTERVAL, values=xs)
    ymin, ymax = domain_axis_bounds(*_padded(min(ys), max(ys)), domain=DOMAIN_HOLD, values=ys)
    body = _axis(xmin, xmax, ymin, ymax, _rt("report_axis_interval_ms"), _rt("report_axis_hold_ms"))
    if reference:
        body += _svg_range_band_vertical(reference, xmin, xmax)
        body += _svg_reference_vertical(reference, xmin, xmax)
    contract = _range_contract(reference)
    for x, y in zip(xs, ys):
        cx = _scale(x, xmin, xmax, _L, _W - _R)
        cy = _scale(y, ymin, ymax, _H - _B, _T)
        if contract and classify_interval(x, contract) in (POSITION_BELOW, POSITION_ABOVE):
            body += _svg_diamond(cx, cy)
        else:
            body += f"<circle cx='{cx:.1f}' cy='{cy:.1f}' r='2.2' class='dot'/>"
    return _svg(body, title)


def _svg_hold_series(series: list[tuple[float, float]]) -> str:
    title = _rt("report_chart_hold_over_time")
    pts = _downsample(series)
    if not pts:
        return _empty_chart(title)
    xs = [p[0] for p in pts]
    ys = [p[1] for p in pts]
    ymin, ymax = domain_axis_bounds(*_padded(min(ys), max(ys)), domain=DOMAIN_HOLD, values=ys)
    xmin, xmax = min(xs), max(xs)
    if xmax <= xmin:
        xmax = xmin + 1.0
    span_ms = max(xmax - xmin, abs(xmax), abs(xmin))
    x_label = _rt("report_axis_elapsed_s") if elapsed_axis_uses_seconds(span_ms) else _rt("report_axis_elapsed")
    body = _axis(
        xmin, xmax, ymin, ymax, x_label, _rt("report_axis_hold_ms"),
        format_x=lambda v, span=span_ms: format_elapsed_axis_tick(v, span_ms=span),
    )
    mapped = [
        (_scale(x, xmin, xmax, _L, _W - _R), _scale(y, ymin, ymax, _H - _B, _T))
        for x, y in zip(xs, ys)
    ]
    if len(mapped) >= 2:
        coords = " ".join(f"{cx:.1f},{cy:.1f}" for cx, cy in mapped)
        body += f"<polyline points='{coords}' class='series'><title>{_rt("report_legend_host_observation")}</title></polyline>"
    if len(mapped) <= 80:
        body += "".join(f"<circle cx='{cx:.1f}' cy='{cy:.1f}' r='2.2' class='dot'/>" for cx, cy in mapped)
    return _svg(body, title)


def _chart_legend(reference: dict[str, Any] | None, host_swatch: str, *, markers: bool = True) -> str:
    if not reference:
        return ""
    items = [
        f"<span class='li'><i class='sw {host_swatch}'></i>{_rt("report_legend_host_observation")}</span>",
    ]
    if _configured_bounds(reference):
        items.append(
            f"<span class='li'><i class='sw sw-band'></i>{_rt("report_legend_configured_range")}</span>"
        )
    if markers and _configured_bounds(reference):
        items.append(
            f"<span class='li'><i class='sw sw-dot'></i>{_rt("report_legend_within_configured_range")}</span>"
        )
        items.append(
            f"<span class='li'><i class='sw sw-out'></i>{_rt("report_legend_outside_configured_range")}</span>"
        )
    items.append(
        f"<span class='li'><i class='sw sw-ref'></i>{_rt("report_legend_configured_reference")} {_e(_reference_label(reference))}</span>"
    )
    return f"<div class='legend'>{''.join(items)}</div>"


_CSS = """
:root{--fg:#1f2933;--muted:#5f6b7a;--line:#d9dee5;--card:#fff;--bg:#f4f6f8;--accent:#2563eb}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--fg);font:14px/1.5 "Segoe UI",system-ui,-apple-system,Arial,sans-serif}
main{max-width:1100px;margin:0 auto;padding:24px 16px 48px}
h1{font-size:24px;margin:0 0 16px}
h2{font-size:17px;margin:0 0 12px}
h3{font-size:14px;margin:0 0 8px;color:var(--muted);text-transform:uppercase;letter-spacing:.04em}
section{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:16px 18px;margin:0 0 16px;min-width:0}
.badge{display:inline-block;font-weight:700;padding:1px 12px;border-radius:999px;border:1px solid transparent}
.v-COMPLETE{background:#dbeafe;color:#1e3a8a;border-color:#93c5fd}
.v-NO_DATA{background:#e5e7eb;color:#374151;border-color:#d1d5db}
.v-INVALID{background:#ffedd5;color:#7c2d12;border-color:#fdba74}
.code{font:13px/1.6 Consolas,"Cascadia Mono",monospace}
table{border-collapse:collapse;width:100%}
th,td{text-align:left;padding:6px 8px;border-bottom:1px solid var(--line);vertical-align:top;overflow-wrap:break-word}
td{overflow-wrap:anywhere}
table.kv th{width:38%;font-weight:600;color:var(--muted)}
.note{font-size:12px;color:var(--muted);margin:0 0 8px}
.cols{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,420px),1fr));gap:16px}
.cols section{margin:0}
.stack{margin:0 0 16px}
.groups{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,300px),1fr));gap:12px}
.groups table.kv th{width:auto}
.groups table.kv td{white-space:nowrap;text-align:right}
.groups.host-meta table.kv td{white-space:normal;text-align:left;overflow-wrap:anywhere}
.charts{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,420px),1fr));gap:12px;margin-top:12px}
figure{margin:0;border:1px solid var(--line);border-radius:8px;padding:8px;min-width:0}
figcaption{font-weight:600;margin-bottom:4px}
svg{display:block;width:100%;height:auto}
svg .ax{stroke:#9aa5b1;stroke-width:1}
svg .tk{font-size:13px;fill:#5f6b7a}
svg .lb{font-size:14px;fill:#3e4c59}
svg .series{fill:none;stroke:var(--accent);stroke-width:1.4}
svg .ref{stroke:#16a34a;stroke-width:1.4;stroke-dasharray:6 4}
svg .rl{font-size:12px;font-weight:600;fill:#15803d;paint-order:stroke;stroke:#fff;stroke-width:3px}
svg .band{fill:#6f9a7b;fill-opacity:.14}
svg .bnd{stroke:#7f9c88;stroke-width:1}
svg .bar{fill:#60a5fa}
svg .dot{fill:#2563eb;fill-opacity:.55}
svg .out{fill:#dc2626}
.legend{display:flex;flex-wrap:wrap;gap:4px 18px;margin-top:6px;font-size:12px;color:var(--muted)}
.legend .li{display:inline-flex;align-items:center;gap:6px;white-space:nowrap}
.sw{display:inline-block;flex:none;width:18px;height:10px}
.sw-line{height:0;border-top:2px solid var(--accent)}
.sw-bar{background:#60a5fa}
.sw-dot{width:8px;height:8px;border-radius:50%;background:#2563eb;opacity:.7}
.sw-out{width:8px;height:8px;background:#dc2626;transform:rotate(45deg)}
.sw-ref{height:0;border-top:2px dashed #16a34a}
.sw-band{background:rgba(111,154,123,.22);border-left:1px solid #7f9c88;border-right:1px solid #7f9c88}
@media (max-width:640px){main{padding:12px 8px 32px}section{padding:12px}table.kv th{width:45%}}
"""


def _run_status_table(view: dict[str, Any]) -> str:
    rows = [
        f"<tr><th scope='row'>{_e(label)}</th><td><span class='badge v-{_e(value)}'>{_e(value)}</span></td></tr>"
        for _row_id, label, value in view["run_status_rows"]
    ]
    return f"<table class='kv'>{''.join(rows)}</table>"


def render_report_html(model: dict[str, Any], view: dict[str, Any]) -> str:
    analysis = model["analysis"]
    reference = view["chart_reference"]
    n = int(view.get("interval_observation_count") or 0)
    charts = (
        "<div class='charts'>"
        f"<figure><figcaption>{_rt("report_chart_interval_over_time")}</figcaption>{_svg_interval_series(analysis['interval_series'], reference)}"
        f"{_chart_legend(reference, 'sw-line')}</figure>"
        f"<figure><figcaption>{_rt("report_hist_press_interval_distribution", n=n)}</figcaption>"
        f"<p class='note'>{_e(_rt("report_hist_note"))}</p>"
        f"{_svg_histogram(analysis['histogram'], reference)}"
        f"{_chart_legend(reference, 'sw-bar', markers=False)}</figure>"
        f"<figure><figcaption>{_rt("report_chart_hold_over_time")}</figcaption>"
        f"{_svg_hold_series(analysis.get('hold_series') or [])}</figure>"
        f"<figure><figcaption>{_rt("report_chart_hold_vs_interval")}</figcaption>{_svg_scatter(analysis['scatter_points'], reference)}"
        f"{_chart_legend(reference, 'sw-dot')}</figure>"
        "</div>"
    )
    host_evidence = ""
    if view["host_evidence"]:
        host_evidence = (
            "<section id='host-evidence'><h2>" + _e(_rt("report_section_host_evidence")) + "</h2><div class='groups'>"
            + "".join(f"<div><h3>{_e(name)}</h3>{_kv_table(rows)}</div>" for name, rows in view["host_evidence"])
            + "</div></section>"
        )
    host_info = ""
    if view["host_information_groups"] or view["host_information"]:
        groups = view["host_information_groups"] or [(_rt("report_section_host_information"), view["host_information"])]
        host_body = "".join(f"<div><h3>{_e(name)}</h3>{_kv_table(rows)}</div>" for name, rows in groups)
        host_info = (
            f"<section id='host-information'><h2>{_e(_rt('report_section_host_information'))}</h2>"
            f"<div class='groups host-meta'>{host_body}</div></section>"
        )
    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{_e(view['title'])} · {_e(model['run'].get('run_id'))}</title>
<style>{_CSS}</style>
</head>
<body>
<main>
<h1>{_e(view['title'])}</h1>
<section id="run-status"><h2>{_e(_rt("report_section_run_status"))}</h2>{_run_status_table(view)}</section>
<div class="cols stack">
<section id="test-information"><h2>{_e(_rt("report_section_test_information"))}</h2>{_kv_table(view['test_information'])}</section>
{host_info}
</div>
<section id="measurements"><h2>{_e(_rt("report_section_measurements"))}</h2>{_kv_table(view['measurements'])}{charts}</section>
{host_evidence}
</main>
</body>
</html>
"""


def write_report_html(
    path: Path,
    model: dict[str, Any],
    *,
    host_info: list[tuple[str, str, str]] | None = None,
    generated_at: str = "",
) -> None:
    view = build_report_view(model, host_info=host_info, generated_at=generated_at)
    path.write_text(render_report_html(model, view), encoding="utf-8")


# --------------------------------------------------------------------------- XLSX

_XLSX_HELPER_COL = 27  # AA
_XLSX_HELPER_SPAN = 5
_XLSX_HISTOGRAM_MAX_BINS = 48


def _xlsx_histogram_bins(
    bins: list[dict[str, Any]], reference: dict[str, Any] | None,
) -> list[tuple[float, float, int]]:
    """Histogram rows for the workbook, padded with empty equal-width bins so the reference span is on-axis."""
    out = [(float(b["lo"]), float(b["hi"]), int(b["count"])) for b in bins]
    span = _reference_span(reference)
    if not out or not span:
        return out
    width = (out[-1][1] - out[0][0]) / len(out)
    if width <= 0:
        return out
    before = max(0, math.ceil((out[0][0] - span[0]) / width - 1e-9))
    after = max(0, math.ceil((span[1] - out[-1][1]) / width - 1e-9))
    if len(out) + before + after > _XLSX_HISTOGRAM_MAX_BINS:
        return out
    lo0, hi_last = out[0][0], out[-1][1]
    head = [(lo0 - (i + 1) * width, lo0 - i * width, 0) for i in reversed(range(before))]
    tail = [(hi_last + i * width, hi_last + (i + 1) * width, 0) for i in range(after)]
    return head + out + tail


def write_report_xlsx(
    path: Path,
    model: dict[str, Any],
    *,
    host_info: list[tuple[str, str, str]] | None = None,
    generated_at: str = "",
) -> bool:
    """Write the V2 workbook. Returns False only when openpyxl is unavailable."""
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.chart import BarChart, Reference, ScatterChart, Series
        from openpyxl.chart.legend import Legend
        from openpyxl.chart.series import SeriesLabel
        from openpyxl.chart.shapes import GraphicalProperties
        from openpyxl.drawing.line import LineProperties
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception:
        return False

    view = build_report_view(model, host_info=host_info, generated_at=generated_at)
    bold = Font(bold=True)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E78")
    section_font = Font(bold=True, size=12, color="1F4E78")

    wb = Workbook()
    ws = wb.active
    sheets = report_sheet_names(model)
    ws.title = sheets[0]

    def kv(sheet, row: int, pairs: list[tuple[str, str, str]]) -> int:
        for row_id, label, value in pairs:
            sheet.cell(row=row, column=1, value=label).font = bold
            cell = sheet.cell(row=row, column=2, value=xlsx_kv_cell_value(row_id, value))
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
            row += 1
        return row

    def section(sheet, row: int, title: str) -> int:
        sheet.cell(row=row, column=1, value=title).font = section_font
        return row + 1

    ws["A1"] = view["title"]
    ws["A1"].font = Font(bold=True, size=16)
    row = section(ws, 3, _rt("report_section_run_status"))
    row = kv(ws, row, view["run_status_rows"])
    status_cell = ws.cell(row=4, column=2)
    status_cell.font = Font(bold=True, size=13)
    status_cell.fill = PatternFill("solid", fgColor=_RUN_STATUS_FILLS.get(view["run_status"], "FFFFFF"))
    row = section(ws, row + 1, _rt("report_section_test_information"))
    row = kv(ws, row, view["test_information"])
    if view["host_information_groups"] or view["host_information"]:
        row = section(ws, row + 1, _rt("report_section_host_information"))
        groups = view["host_information_groups"] or [(_rt("report_section_host_information"), view["host_information"])]
        for name, rows in groups:
            ws.cell(row=row, column=1, value=name).font = bold
            row += 1
            row = kv(ws, row, rows) + 1
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 48

    wm = wb.create_sheet(sheets[1])
    wm["A1"] = _rt("report_section_measurements")
    wm["A1"].font = Font(bold=True, size=14)
    row = kv(wm, 3, view["measurements"])
    n_intervals = int(view.get("interval_observation_count") or 0)
    hist_row = section(wm, row + 1, _rt("report_hist_press_interval_distribution", n=n_intervals))
    wm.cell(row=hist_row, column=1, value=_rt("report_hist_note")).alignment = Alignment(wrap_text=True)
    wm.merge_cells(start_row=hist_row, start_column=1, end_row=hist_row, end_column=4)
    hist_row += 1
    for c, name in enumerate((
        _rt("report_hist_bin_centre"),
        _rt("report_hist_low"),
        _rt("report_hist_high"),
        _rt("report_hist_count"),
    ), start=1):
        h = wm.cell(row=hist_row, column=c, value=name)
        h.font = head_font
        h.fill = head_fill
    reference = view["chart_reference"]
    bins = _xlsx_histogram_bins(model["analysis"]["histogram"], reference)
    for i, (lo, hi, count) in enumerate(bins, start=1):
        wm.cell(row=hist_row + i, column=1, value=round((lo + hi) / 2, 1)).number_format = "0.0"
        wm.cell(row=hist_row + i, column=2, value=round(lo, 3))
        wm.cell(row=hist_row + i, column=3, value=round(hi, 3))
        wm.cell(row=hist_row + i, column=4, value=count)
    wm.column_dimensions["A"].width = 34
    wm.column_dimensions["B"].width = 18
    wm.column_dimensions["C"].width = 12
    wm.column_dimensions["D"].width = 10

    if view["host_evidence"]:
        wh = wb.create_sheet(_rt("report_sheet_host_evidence"))
        wh["A1"] = _rt("report_section_host_evidence")
        wh["A1"].font = Font(bold=True, size=14)
        row = 3
        for name, rows in view["host_evidence"]:
            row = section(wh, row, name)
            row = kv(wh, row, rows) + 1
        wh.column_dimensions["A"].width = 52
        wh.column_dimensions["B"].width = 16

    wc = wb.create_sheet(_rt("report_sheet_click_data"))
    cols = view["click_columns"]
    for c, name in enumerate(cols, start=1):
        h = wc.cell(row=1, column=c, value=name)
        h.font = head_font
        h.fill = head_fill
    elapsed_col = 2
    range_col = len(cols) if view.get("click_range_positions") else None
    outside_labels = {_rt("report_value_range_below"), _rt("report_value_range_above")}
    outside_font = Font(color="C00000")
    for r, values in enumerate(view["click_rows"], start=2):
        for c, value in enumerate(values, start=1):
            if c == elapsed_col:
                days = excel_elapsed_days(value)
                cell = wc.cell(row=r, column=c, value=days)
                if days is not None:
                    cell.number_format = EXCEL_ELAPSED_NUMBER_FORMAT
                continue
            cell = wc.cell(row=r, column=c, value=value)
            if isinstance(value, float):
                cell.number_format = "0.000"
            if range_col is not None and c == range_col and str(value) in outside_labels:
                cell.font = outside_font
    n_rows = len(view["click_rows"])
    wc.freeze_panes = "A2"
    wc.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n_rows + 1}"
    for c, width in enumerate((8, 14, 10, 14, 12, 18, 34, 20, 28)[: len(cols)], start=1):
        wc.column_dimensions[get_column_letter(c)].width = width

    def frame(chart, title: str, x_title: str, y_title: str, x_range=None, y_range=None):
        chart.title = title
        chart.x_axis.title = x_title
        chart.y_axis.title = y_title
        for item in (chart.title, chart.x_axis.title, chart.y_axis.title):
            item.overlay = False
        # openpyxl >= 3.1 otherwise writes axes that Excel hides.
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.x_axis.number_format = "0"
        chart.y_axis.number_format = "0"
        for axis, bounds in ((chart.x_axis, x_range), (chart.y_axis, y_range)):
            if bounds:
                axis.scaling.min, axis.scaling.max = (math.floor(bounds[0]), math.ceil(bounds[1]))
        chart.varyColors = False
        chart.legend = None
        chart.height, chart.width = 7.5, 16
        return chart

    # Reference marker coordinates live in hidden helper columns of the Measurements sheet.
    helper_row = [0]

    def marker_series(xs: list[float | None], ys: list[float | None], title: str, dash: str, color: str = "16A34A"):
        helper_row[0] += 1
        r = helper_row[0]
        for i, (x, y) in enumerate(zip(xs, ys)):
            wm.cell(row=r, column=_XLSX_HELPER_COL + i, value=x)
            wm.cell(row=r, column=_XLSX_HELPER_COL + _XLSX_HELPER_SPAN + i, value=y)
        n = len(xs)
        s = Series(
            Reference(wm, min_col=_XLSX_HELPER_COL + _XLSX_HELPER_SPAN, max_col=_XLSX_HELPER_COL + _XLSX_HELPER_SPAN + n - 1, min_row=r),
            Reference(wm, min_col=_XLSX_HELPER_COL, max_col=_XLSX_HELPER_COL + n - 1, min_row=r),
            title=title,
        )
        s.marker.symbol = "none"
        s.smooth = False
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 15875
        s.graphicalProperties.line.dashStyle = dash
        return s

    def bound_markers(vertical: bool, lo: float, hi: float) -> list:
        bounds = _configured_bounds(reference)
        if not bounds:
            return []
        lower, upper = bounds
        if vertical:
            return [
                marker_series([lower, lower], [lo, hi], f"{_rt("report_label_configured_interval_range")} {lower:.2f} ms", "dash", "64748B"),
                marker_series([upper, upper], [lo, hi], f"{_rt("report_label_configured_interval_range")} {upper:.2f} ms", "dash", "64748B"),
            ]
        return [
            marker_series([lo, hi], [lower, lower], f"{_rt("report_label_configured_interval_range")} {lower:.2f} ms", "dash", "64748B"),
            marker_series([lo, hi], [upper, upper], f"{_rt("report_label_configured_interval_range")} {upper:.2f} ms", "dash", "64748B"),
        ]

    def reference_markers(vertical: bool, lo: float, hi: float) -> list:
        ref = reference["interval_ms"]
        if vertical:
            return [marker_series([ref, ref], [lo, hi], f"{_rt("report_legend_configured_reference")} {ref:.2f} ms", "dash")]
        return [marker_series([lo, hi], [ref, ref], f"{_rt("report_legend_configured_reference")} {ref:.2f} ms", "dash")]

    def with_legend(chart):
        chart.legend = Legend()
        chart.legend.position = "b"
        chart.legend.overlay = False
        chart.visible_cells_only = False
        chart.height = 8.5
        return chart

    span = _reference_span(reference)
    series = model["analysis"]["interval_series"]
    points = model["analysis"]["scatter_points"]
    if n_rows >= 2 and series and points:
        last = n_rows + 1
        ys = [v for _, v in series] + (list(span) if span else [])
        line = frame(
            ScatterChart(scatterStyle="line"), _rt("report_chart_interval_over_time"), _rt("report_axis_elapsed"), _rt("report_axis_interval_ms"),
            y_range=domain_axis_bounds(*_padded(min(ys), max(ys)), domain=DOMAIN_INTERVAL, values=[v for _, v in series]),
        )
        line.x_axis.number_format = "[h]:mm:ss"
        s = Series(Reference(wc, min_col=4, min_row=2, max_row=last), Reference(wc, min_col=2, min_row=2, max_row=last),
                   title=_rt("report_legend_host_observation"))
        s.marker.symbol = "none"
        s.smooth = False
        s.graphicalProperties.line.solidFill = "2563EB"
        s.graphicalProperties.line.width = 15875
        line.series.append(s)
        if reference:
            x0 = excel_elapsed_days(series[0][0])
            x1 = excel_elapsed_days(series[-1][0])
            x_lo = x0 if x0 is not None else 0.0
            x_hi = x1 if x1 is not None else 0.0
            line.series.extend(bound_markers(False, x_lo, x_hi))
            line.series.extend(reference_markers(False, x_lo, x_hi))
            with_legend(line)
        wm.add_chart(line, "F3")

        xs = [p[0] for p in points] + (list(span) if span else [])
        holds = [p[1] for p in points]
        y_range = domain_axis_bounds(*_padded(min(holds), max(holds)), domain=DOMAIN_HOLD, values=holds)
        x_range = domain_axis_bounds(*_padded(min(xs), max(xs)), domain=DOMAIN_INTERVAL, values=[p[0] for p in points])
        scatter = frame(
            ScatterChart(scatterStyle="marker"), _rt("report_chart_hold_vs_interval"), _rt("report_axis_interval_ms"), _rt("report_axis_hold_ms"),
            x_range=x_range, y_range=y_range,
        )
        scatter.y_axis.number_format = "0.0"
        p = Series(Reference(wc, min_col=5, min_row=2, max_row=last), Reference(wc, min_col=4, min_row=2, max_row=last),
                   title=_rt("report_legend_host_observation"))
        p.marker.symbol = "circle"
        p.marker.size = 5
        p.marker.graphicalProperties.solidFill = "2563EB"
        p.marker.graphicalProperties.line.solidFill = "2563EB"
        p.graphicalProperties.line.noFill = True
        scatter.series.append(p)
        if reference:
            y_lo, y_hi = math.floor(y_range[0]), math.ceil(y_range[1])
            scatter.series.extend(bound_markers(True, y_lo, y_hi))
            scatter.series.extend(reference_markers(True, y_lo, y_hi))
            with_legend(scatter)
        wm.add_chart(scatter, "F37")
    hold_series = model["analysis"].get("hold_series") or []
    if n_rows >= 2 and hold_series:
        last = n_rows + 1
        hold_values = [v for _, v in hold_series]
        hold_line = frame(
            ScatterChart(scatterStyle="line"), _rt("report_chart_hold_over_time"), _rt("report_axis_elapsed"), _rt("report_axis_hold_ms"),
            y_range=domain_axis_bounds(*_padded(min(hold_values), max(hold_values)), domain=DOMAIN_HOLD, values=hold_values),
        )
        hold_line.x_axis.number_format = "[h]:mm:ss"
        hs = Series(Reference(wc, min_col=5, min_row=2, max_row=last), Reference(wc, min_col=2, min_row=2, max_row=last),
                    title=_rt("report_legend_host_observation"))
        hs.marker.symbol = "none"
        hs.smooth = False
        hs.graphicalProperties.line.solidFill = "64748B"
        hs.graphicalProperties.line.width = 15875
        hold_line.series.append(hs)
        wm.add_chart(hold_line, "F54")
    if bins:
        bar = frame(BarChart(), _rt("report_chart_interval_distribution"), _rt("report_axis_interval_ms"), _rt("report_axis_count"))
        bar.x_axis.axPos = "b"
        bar.gapWidth = 30
        bar.add_data(Reference(wm, min_col=4, min_row=hist_row + 1, max_row=hist_row + len(bins)))
        bar.set_categories(Reference(wm, min_col=1, min_row=hist_row + 1, max_row=hist_row + len(bins)))
        bar.series[0].tx = SeriesLabel(v=_rt("report_legend_host_observation"))
        bar.series[0].graphicalProperties.solidFill = "60A5FA"
        bar.series[0].graphicalProperties.line.solidFill = "60A5FA"
        if reference and bins[0][0] <= span[0] and span[1] <= bins[-1][1]:
            # Overlay on its own value axes spanning the bin edges so markers sit at exact interval values.
            overlay = ScatterChart()
            overlay.x_axis.axId, overlay.y_axis.axId = 500, 600
            overlay.x_axis.crossAx, overlay.y_axis.crossAx = 600, 500
            overlay.x_axis.axPos, overlay.y_axis.axPos = "t", "r"
            overlay.y_axis.crosses = "max"
            overlay.x_axis.scaling.min, overlay.x_axis.scaling.max = bins[0][0], bins[-1][1]
            overlay.y_axis.scaling.min, overlay.y_axis.scaling.max = 0, 1
            for axis in (overlay.x_axis, overlay.y_axis):
                # A deleted secondary axis makes Excel plot the series on the primary axes.
                axis.delete = False
                axis.majorGridlines = None
                axis.majorTickMark = "none"
                axis.number_format = ";;;"
                axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
            overlay.series.extend(bound_markers(True, 0, 1))
            overlay.series.extend(reference_markers(True, 0, 1))
            bar += overlay
            with_legend(bar)
        wm.add_chart(bar, "F20")
    if helper_row[0]:
        for c in range(_XLSX_HELPER_COL, _XLSX_HELPER_COL + 2 * _XLSX_HELPER_SPAN):
            wm.column_dimensions[get_column_letter(c)].hidden = True

    wb.active = 0
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    load_workbook(tmp).close()
    os.replace(tmp, path)
    return True
